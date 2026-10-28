import datetime
import io
from unittest.mock import patch
from uuid import uuid4

import openpyxl
import pytest
from django.db import connection
from django.http import QueryDict
from django.test.utils import CaptureQueriesContext
from model_bakery import baker
from rest_framework import status

from src.dados_comuns.constants import (
    ADMINISTRADOR_EMPRESA,
    ADMINISTRADOR_GESTAO_ALIMENTACAO_TERCEIRIZADA,
    DIRETOR_UE,
    FORMATO_DATA_BRASILEIRO,
    MENSAGEM_SOLICITACAO_GERACAO_ARQUIVO,
    TIPOS_GESTAO,
)
from src.dados_comuns.models import CentralDeDownload
from src.escola.models import GrupoUnidadeEscolar
from src.medicao_inicial.models import CategoriaMedicao
from src.medicao_inicial.services.relatorio_alimentacoes_servidas import (
    filtros_exportacao_relatorio_alimentacoes_servidas,
)
from src.medicao_inicial.services.relatorio_alimentacoes_servidas_dados import (
    SECAO_DIETA_TIPO_A,
    SECAO_DIETA_TIPO_B,
    SECAO_REGULAR,
    consolida_secoes,
    monta_colunas,
    monta_relatorio_alimentacoes_servidas,
    obtem_nome_arquivo_relatorio_alimentacoes_servidas,
    soma_colunas,
)
from src.medicao_inicial.services.relatorio_alimentacoes_servidas_excel import (
    gera_relatorio_alimentacoes_servidas_xlsx,
)
from src.medicao_inicial.utils import calcula_totais_consumo_por_grupo
from src.terceirizada.models import Terceirizada

pytestmark = pytest.mark.django_db

URL = "/medicao-inicial/relatorios/relatorio-alimentacoes-servidas/exportar-xlsx/"
DELAY = (
    "src.medicao_inicial.api.viewsets."
    "exporta_relatorio_alimentacoes_servidas_para_xlsx.delay"
)
DEZEMBRO = 0
MENSAGEM_DATA_FORA = "A data deve estar dentro de um mês de referência selecionado"


def _usuario(django_user_model, instituicao, perfil_nome):
    sufixo = uuid4().hex[:10]
    email = f"{sufixo}@test.com"
    usuario = django_user_model.objects.create_user(
        username=email,
        password=uuid4().hex,
        email=email,
        registro_funcional=sufixo[:7],
    )
    baker.make(
        "Vinculo",
        usuario=usuario,
        instituicao=instituicao,
        perfil=baker.make("Perfil", nome=perfil_nome, ativo=True),
        data_inicial=datetime.date.today(),
        ativo=True,
    )
    return usuario


def _contrato(terceirizada, lote):
    edital = baker.make(
        "Edital",
        numero=f"ED-{uuid4().hex[:8]}",
        tipo_contratacao="PREGAO",
        processo=f"PROC-{uuid4().hex[:8]}",
        objeto="Alimentacao",
    )
    contrato = baker.make(
        "Contrato",
        numero=f"CT-{uuid4().hex[:8]}",
        processo=f"PROC-{uuid4().hex[:8]}",
        terceirizada=terceirizada,
        edital=edital,
        modalidade=baker.make("Modalidade", nome="Pregão Eletrônico"),
        encerrado=False,
    )
    contrato.lotes.add(lote)


def _escola(base, dre, lote, tipo, subprefeitura):
    return baker.make(
        "Escola",
        codigo_eol=f"{uuid4().int % 900000 + 100000:06d}",
        nome=f"UE {tipo.iniciais} {uuid4().hex[:4]}",
        diretoria_regional=dre,
        lote=lote,
        tipo_unidade=tipo,
        tipo_gestao=base["gestao"],
        subprefeitura=subprefeitura,
    )


def _medicao(base, escola, mes="12", ano="2023", recreio=None):
    solicitacao = baker.make(
        "SolicitacaoMedicaoInicial",
        escola=escola,
        mes=mes,
        ano=ano,
        status="MEDICAO_APROVADA_PELA_CODAE",
        recreio_nas_ferias=recreio,
    )
    return baker.make(
        "Medicao",
        solicitacao_medicao_inicial=solicitacao,
        periodo_escolar=base["periodo"],
    )


CAMPO_POR_ALIMENTACAO = {"Lanche": "lanche", "Refeição": "refeicao"}


def _valor(medicao, categoria, valor, dia="01", alimentacao=None, nome_campo=None):
    baker.make(
        "ValorMedicao",
        medicao=medicao,
        categoria_medicao=categoria,
        valor=str(valor),
        dia=dia,
        semana="1",
        nome_campo=nome_campo or CAMPO_POR_ALIMENTACAO[alimentacao.nome],
        tipo_alimentacao=alimentacao,
    )


def _grupo(nome, *tipos):
    grupo = baker.make("GrupoUnidadeEscolar", nome=nome)
    grupo.tipos_unidades.add(*tipos)


def _cadastros(base):
    base["gestao"] = baker.make("TipoGestao", nome=TIPOS_GESTAO.TERC_TOTAL.value)
    base["empresa"] = baker.make("Terceirizada", tipo_servico=Terceirizada.TERCEIRIZADA)
    base["outra_empresa"] = baker.make(
        "Terceirizada", tipo_servico=Terceirizada.TERCEIRIZADA
    )
    base["dre"] = baker.make("DiretoriaRegional", nome="DRE IPIRANGA")
    base["outra_dre"] = baker.make("DiretoriaRegional", nome="DRE PENHA")
    base["lote"] = baker.make(
        "Lote",
        nome="LOTE 1",
        diretoria_regional=base["dre"],
        terceirizada=base["empresa"],
    )
    base["lote_alheio"] = baker.make(
        "Lote",
        nome="LOTE 2",
        diretoria_regional=base["outra_dre"],
        terceirizada=base["outra_empresa"],
    )
    base["lote_outra_empresa"] = baker.make(
        "Lote",
        nome="LOTE 3",
        diretoria_regional=base["dre"],
        terceirizada=base["outra_empresa"],
    )
    _contrato(base["empresa"], base["lote"])
    base["sub"] = baker.make("Subprefeitura", nome="IPIRANGA")
    base["sub"].diretoria_regional.add(base["dre"])
    base["outra_sub"] = baker.make("Subprefeitura", nome="PENHA")
    base["outra_sub"].diretoria_regional.add(base["outra_dre"])
    base["periodo"] = baker.make("PeriodoEscolar", nome="MANHA")


def _tipos(base):
    for chave, iniciais in (
        ("emei", "EMEI"),
        ("ceu_emei", "CEU EMEI"),
        ("cemei", "CEMEI"),
        ("cei", "CEI"),
    ):
        base[f"tipo_{chave}"] = baker.make("TipoUnidadeEscolar", iniciais=iniciais)
    _grupo(GrupoUnidadeEscolar.GRUPO_1, base["tipo_cei"])
    _grupo(GrupoUnidadeEscolar.GRUPO_2, base["tipo_cemei"])
    _grupo(GrupoUnidadeEscolar.GRUPO_3, base["tipo_emei"], base["tipo_ceu_emei"])


def _dominio_medicao(base):
    base["regular"] = baker.make("CategoriaMedicao", nome=CategoriaMedicao.ALIMENTACAO)
    base["tipo_a"] = baker.make(
        "CategoriaMedicao", nome=CategoriaMedicao.DIETA_ESPECIAL_TIPO_A
    )
    base["tipo_a_enteral"] = baker.make(
        "CategoriaMedicao",
        nome=CategoriaMedicao.DIETA_ESPECIAL_TIPO_A_ENTERAL_RESTRICAO_AMINOACIDOS,
    )
    base["tipo_b"] = baker.make(
        "CategoriaMedicao", nome=CategoriaMedicao.DIETA_ESPECIAL_TIPO_B
    )
    base["lanche"] = baker.make("TipoAlimentacao", nome="Lanche")
    base["refeicao"] = baker.make("TipoAlimentacao", nome="Refeição")


def _escolas(base):
    base["emei"] = _escola(
        base, base["dre"], base["lote"], base["tipo_emei"], base["sub"]
    )
    base["ceu_emei"] = _escola(
        base, base["dre"], base["lote"], base["tipo_ceu_emei"], base["sub"]
    )
    base["cemei"] = _escola(
        base, base["dre"], base["lote"], base["tipo_cemei"], base["sub"]
    )
    base["emei_alheia"] = _escola(
        base,
        base["outra_dre"],
        base["lote_alheio"],
        base["tipo_emei"],
        base["outra_sub"],
    )
    base["emei_outra_empresa"] = _escola(
        base, base["dre"], base["lote_outra_empresa"], base["tipo_emei"], base["sub"]
    )


def _lancamentos_dezembro(base):
    medicao = _medicao(base, base["emei"])
    _valor(medicao, base["regular"], 10, alimentacao=base["lanche"])
    _valor(medicao, base["regular"], 5, dia="20", alimentacao=base["lanche"])
    _valor(medicao, base["regular"], 7, alimentacao=base["refeicao"])
    _valor(medicao, base["regular"], 100, nome_campo="matriculados")
    _valor(medicao, base["tipo_a"], 2, alimentacao=base["lanche"])
    _valor(medicao, base["tipo_a_enteral"], 1, alimentacao=base["refeicao"])
    _valor(medicao, base["tipo_b"], 3, alimentacao=base["lanche"])
    _valor(
        _medicao(base, base["ceu_emei"]), base["regular"], 4, alimentacao=base["lanche"]
    )
    _valor(
        _medicao(base, base["emei_alheia"]),
        base["regular"],
        8,
        alimentacao=base["lanche"],
    )
    _valor(
        _medicao(base, base["emei_outra_empresa"]),
        base["regular"],
        50,
        alimentacao=base["lanche"],
    )
    cemei = _medicao(base, base["cemei"])
    _valor(cemei, base["regular"], 6, alimentacao=base["lanche"])
    _valor(cemei, base["tipo_a"], 2, alimentacao=base["lanche"])


def _lancamentos_outubro_e_novembro(base):
    outubro = _medicao(base, base["emei"], mes="10", ano="2023")
    _valor(outubro, base["regular"], 3, dia="05", alimentacao=base["lanche"])
    _valor(outubro, base["regular"], 4, dia="15", alimentacao=base["lanche"])
    novembro = _medicao(base, base["emei"], mes="11", ano="2023")
    _valor(novembro, base["regular"], 100, dia="10", alimentacao=base["lanche"])


def _lancamentos_janeiro(base):
    _valor(
        _medicao(base, base["emei"], mes="01", ano="2024"),
        base["regular"],
        11,
        alimentacao=base["lanche"],
    )
    recreio = baker.make(
        "RecreioNasFerias",
        titulo="Recreio",
        data_inicio=datetime.date(2024, 1, 2),
        data_fim=datetime.date(2024, 1, 31),
    )
    base["recreio"] = recreio
    _valor(
        _medicao(base, base["emei"], mes="01", ano="2024", recreio=recreio),
        base["regular"],
        9,
        alimentacao=base["lanche"],
    )


@pytest.fixture
def base(django_user_model):
    dados = {}
    _cadastros(dados)
    _tipos(dados)
    _dominio_medicao(dados)
    _escolas(dados)
    _lancamentos_dezembro(dados)
    _lancamentos_outubro_e_novembro(dados)
    _lancamentos_janeiro(dados)
    dados["usuario"] = _usuario(
        django_user_model,
        baker.make("Codae"),
        ADMINISTRADOR_GESTAO_ALIMENTACAO_TERCEIRIZADA,
    )
    return dados


def _params(base, **extras):
    params = {"meses[]": ["12_2023"], "dres[]": [str(base["dre"].uuid)]}
    params.update(extras)
    return params


def _uuids(*objetos):
    return [str(objeto.uuid) for objeto in objetos]


def _relatorio(base, usuario=None, **extras):
    query = QueryDict(mutable=True)
    for chave, valor in _params(base, **extras).items():
        if isinstance(valor, list):
            query.setlist(chave, valor)
        else:
            query[chave] = valor
    usuario = usuario or base["usuario"]
    filtros = filtros_exportacao_relatorio_alimentacoes_servidas(query, usuario)
    return monta_relatorio_alimentacoes_servidas(filtros, usuario)


def _grupo_do_relatorio(relatorio, nome):
    return next(grupo for grupo in relatorio["grupos"] if grupo["nome"] == nome)


def _textos(arquivo):
    return [
        celula for linha in _linhas(arquivo) for celula in linha if celula is not None
    ]


def _linhas(arquivo):
    """Linhas de todas as abas, na ordem das abas."""
    workbook = openpyxl.load_workbook(io.BytesIO(arquivo))
    linhas = [
        linha
        for aba in workbook.worksheets
        for linha in aba.iter_rows(values_only=True)
    ]
    workbook.close()
    return linhas


def test_exportar_sem_meses_ou_sem_dres_retorna_400(client, base):
    client.force_login(base["usuario"])

    sem_meses = client.get(URL, {"dres[]": _uuids(base["dre"])})
    assert sem_meses.status_code == status.HTTP_400_BAD_REQUEST
    assert sem_meses.json()["meses"] == ["Selecione ao menos um mês de referência."]

    sem_dres = client.get(URL, {"meses[]": ["12_2023"]})
    assert sem_dres.status_code == status.HTTP_400_BAD_REQUEST
    assert sem_dres.json()["dres"] == ["Informe ao menos uma diretoria regional."]


def test_exportar_com_mes_e_dres_registra_na_central_de_downloads(client, base):
    client.force_login(base["usuario"])

    resposta = client.get(URL, _params(base))

    assert resposta.status_code == status.HTTP_200_OK
    assert resposta.json() == {"detail": MENSAGEM_SOLICITACAO_GERACAO_ARQUIVO}
    central = CentralDeDownload.objects.get(usuario=base["usuario"])
    assert central.identificador == "Relatório de Alimentações Servidas - 12/2023.xlsx"
    assert central.status == CentralDeDownload.STATUS_CONCLUIDO
    with central.arquivo.open("rb") as arquivo:
        textos = _textos(arquivo.read())
    assert "Relatório de Alimentações Servidas" in textos
    assert any("DRE IPIRANGA" in str(texto) for texto in textos)
    hoje = datetime.date.today().strftime(FORMATO_DATA_BRASILEIRO)
    assert f"Data: {hoje}" in textos
    assert not any("MÉDIA" in str(texto).upper() for texto in textos)


def test_exportar_le_listas_repetidas_de_dres_e_lotes(client, base):
    client.force_login(base["usuario"])

    with patch(DELAY) as delay:
        resposta = client.get(
            URL,
            _params(
                base,
                **{
                    "dres[]": _uuids(base["dre"], base["outra_dre"]),
                    "lotes[]": _uuids(base["lote"], base["lote_alheio"]),
                },
            ),
        )

    assert resposta.status_code == status.HTTP_200_OK
    filtros = delay.call_args.kwargs["filtros"]
    assert set(filtros["dres"]) == set(_uuids(base["dre"], base["outra_dre"]))
    assert set(filtros["lotes"]) == set(_uuids(base["lote"], base["lote_alheio"]))
    relatorio = monta_relatorio_alimentacoes_servidas(filtros, base["usuario"])
    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    assert (
        grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"][DEZEMBRO] == 10 + 5 + 4 + 8
    )
    assert relatorio["dres"] == ["DRE IPIRANGA", "DRE PENHA"]


def test_parametro_omitido_nao_zera_e_preenchido_restringe(base):
    completo = _relatorio(base)
    grupo_3 = _grupo_do_relatorio(completo, GrupoUnidadeEscolar.GRUPO_3)
    assert (
        grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"][DEZEMBRO] == 10 + 5 + 4 + 50
    )
    assert {grupo["nome"] for grupo in completo["grupos"]} == {
        GrupoUnidadeEscolar.GRUPO_2,
        GrupoUnidadeEscolar.GRUPO_3,
    }

    por_lote = _relatorio(base, **{"lotes[]": _uuids(base["lote"])})
    grupo_3 = _grupo_do_relatorio(por_lote, GrupoUnidadeEscolar.GRUPO_3)
    assert grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"][DEZEMBRO] == 10 + 5 + 4


def test_filtro_por_subprefeitura(base):
    relatorio = _relatorio(
        base,
        **{
            "dres[]": _uuids(base["dre"], base["outra_dre"]),
            "subprefeituras[]": _uuids(base["outra_sub"]),
        },
    )

    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    assert list(grupo_3["alimentacoes"][SECAO_REGULAR]) == ["Lanche 5h"]
    assert grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"][DEZEMBRO] == 8
    assert [grupo["nome"] for grupo in relatorio["grupos"]] == [
        GrupoUnidadeEscolar.GRUPO_3
    ]


def test_exportar_rejeita_lote_junto_com_subprefeitura(client, base):
    client.force_login(base["usuario"])

    resposta = client.get(
        URL,
        _params(
            base,
            **{
                "lotes[]": _uuids(base["lote"]),
                "subprefeituras[]": _uuids(base["sub"]),
            },
        ),
    )

    assert resposta.status_code == status.HTTP_400_BAD_REQUEST
    assert resposta.json()["subprefeituras"] == [
        "Subprefeitura não pode ser informada junto com lote."
    ]


def test_tipos_unidades_parcial_nao_reconstroi_o_grupo(base):
    relatorio = _relatorio(base, **{"tipos_unidades[]": _uuids(base["tipo_emei"])})

    assert [grupo["nome"] for grupo in relatorio["grupos"]] == [
        GrupoUnidadeEscolar.GRUPO_3
    ]
    grupo_3 = relatorio["grupos"][0]
    assert grupo_3["tipos_unidades"] == ["EMEI"]
    assert grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"][DEZEMBRO] == 10 + 5 + 50


def test_unidade_tipo_de_alimentacao_e_dietas(base):
    relatorio = _relatorio(
        base,
        **{
            "unidades_educacionais[]": _uuids(base["emei"]),
            "tipos_alimentacao[]": _uuids(base["refeicao"]),
        },
    )

    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    assert list(grupo_3["alimentacoes"][SECAO_REGULAR]) == ["Refeição"]
    assert grupo_3["alimentacoes"][SECAO_REGULAR]["Refeição"][DEZEMBRO] == 7
    assert grupo_3["alimentacoes"][SECAO_DIETA_TIPO_A]["Refeição"][DEZEMBRO] == 1
    assert grupo_3["alimentacoes"][SECAO_DIETA_TIPO_B] == {}

    completo = _relatorio(base, **{"unidades_educacionais[]": _uuids(base["emei"])})
    grupo_3 = _grupo_do_relatorio(completo, GrupoUnidadeEscolar.GRUPO_3)
    assert grupo_3["alimentacoes"][SECAO_DIETA_TIPO_A]["Lanche 5h"][DEZEMBRO] == 2
    assert grupo_3["alimentacoes"][SECAO_DIETA_TIPO_B]["Lanche 5h"][DEZEMBRO] == 3


def test_periodo_limita_os_dias(client, base):
    relatorio = _relatorio(
        base,
        periodo_lancamento_de="01/12/2023",
        periodo_lancamento_ate="10/12/2023",
        **{"unidades_educacionais[]": _uuids(base["emei"])},
    )
    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    assert grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"][DEZEMBRO] == 10

    client.force_login(base["usuario"])
    fora_do_mes = client.get(
        URL,
        _params(
            base,
            periodo_lancamento_de="01/11/2023",
            periodo_lancamento_ate="10/11/2023",
        ),
    )
    assert fora_do_mes.status_code == status.HTTP_400_BAD_REQUEST
    assert fora_do_mes.json()["periodo_lancamento_de"] == [MENSAGEM_DATA_FORA]

    invertido = client.get(
        URL,
        _params(
            base,
            periodo_lancamento_de="20/12/2023",
            periodo_lancamento_ate="10/12/2023",
        ),
    )
    assert invertido.status_code == status.HTTP_400_BAD_REQUEST
    assert "anterior" in invertido.json()["periodo_lancamento_de"][0]


def test_perfil_autorizado_e_escola_negada(client, base, django_user_model):
    usuario_dre = _usuario(django_user_model, base["dre"], "COGESTOR_DRE")
    client.force_login(usuario_dre)
    with patch(DELAY):
        assert client.get(URL, _params(base)).status_code == status.HTTP_200_OK

    usuario_escola = _usuario(django_user_model, base["emei"], DIRETOR_UE)
    client.force_login(usuario_escola)
    assert client.get(URL, _params(base)).status_code == status.HTTP_403_FORBIDDEN


def test_dre_lote_e_unidade_fora_do_escopo(client, base, django_user_model):
    usuario_dre = _usuario(django_user_model, base["dre"], "COGESTOR_DRE")
    client.force_login(usuario_dre)

    dre_fora = client.get(URL, _params(base, **{"dres[]": _uuids(base["outra_dre"])}))
    assert dre_fora.status_code == status.HTTP_400_BAD_REQUEST
    assert "dres" in dre_fora.json()

    lote_fora = client.get(
        URL, _params(base, **{"lotes[]": _uuids(base["lote_alheio"])})
    )
    assert lote_fora.status_code == status.HTTP_400_BAD_REQUEST
    assert "lotes" in lote_fora.json()

    unidade_fora = client.get(
        URL, _params(base, **{"unidades_educacionais[]": _uuids(base["emei_alheia"])})
    )
    assert unidade_fora.status_code == status.HTTP_400_BAD_REQUEST
    assert "unidades_educacionais" in unidade_fora.json()


def test_terceirizada_so_recebe_os_proprios_lotes(base, django_user_model):
    usuario = _usuario(django_user_model, base["empresa"], ADMINISTRADOR_EMPRESA)

    relatorio = _relatorio(base, usuario=usuario)

    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    assert grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"][DEZEMBRO] == 10 + 5 + 4


def test_grupo_1_rejeita_tipo_de_alimentacao(client, base):
    client.force_login(base["usuario"])

    grupo_1 = client.get(
        URL,
        _params(
            base,
            **{
                "tipos_unidades[]": _uuids(base["tipo_cei"]),
                "tipos_alimentacao[]": _uuids(base["lanche"]),
            },
        ),
    )
    assert grupo_1.status_code == status.HTTP_400_BAD_REQUEST
    assert "Grupo 1" in grupo_1.json()["tipos_alimentacao"][0]


def _consolidado(relatorio):
    """Soma, por seção e rótulo, as linhas de todos os grupos."""
    return consolida_secoes(
        relatorio["grupos"], "alimentacoes", len(relatorio["colunas"])
    )


def test_excel_tem_regulares_dietas_e_total_geral(base):
    arquivo = gera_relatorio_alimentacoes_servidas_xlsx(_relatorio(base))
    textos = _textos(arquivo)

    for secao in (SECAO_REGULAR, SECAO_DIETA_TIPO_A, SECAO_DIETA_TIPO_B):
        assert secao in textos
    assert "EMEI, CEU EMEI" in textos
    assert not any("GRUPO 3" in str(texto).upper() for texto in textos)
    assert "CONSOLIDAÇÃO - TOTAL GERAL" in textos
    assert not any("TIPOS DE ALIMENTAÇÃO" in str(texto).upper() for texto in textos)
    assert not any("FAIXA" in str(texto).upper() for texto in textos)


def _com_recreio(base, meses, **extras):
    return _relatorio(
        base,
        **{
            "meses[]": meses,
            "recreios_nas_ferias[]": _uuids(base["recreio"]),
            "unidades_educacionais[]": _uuids(base["emei"]),
        },
        **extras,
    )


def test_recreio_selecionado_vai_para_coluna_propria(base):
    relatorio = _com_recreio(base, ["01_2024"])
    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)

    assert grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"] == [11, 9]
    assert relatorio["colunas"] == [
        "REGULAR - JAN 2024",
        "RECREIO NAS FÉRIAS - JAN 2024",
    ]
    assert relatorio["recreios"] == [{"titulo": "Recreio", "meses": [(1, 2024)]}]

    linhas = _linhas(gera_relatorio_alimentacoes_servidas_xlsx(relatorio))
    cabecalho = next(linha for linha in linhas if linha[0] == "Tipo de Alimentação")
    assert cabecalho == (
        "Tipo de Alimentação",
        "REGULAR - JAN 2024",
        "RECREIO NAS FÉRIAS - JAN 2024",
        "TOTAL 2024",
    )
    linha_lanche = next(linha for linha in linhas if linha[0] == "Lanche 5h")
    assert linha_lanche == ("Lanche 5h", 11, 9, 20)


def test_totais_e_consolidacoes_batem_com_o_detalhe(base):
    relatorio = _relatorio(base, **{"tipos_unidades[]": _uuids(base["tipo_cemei"])})
    completo = _relatorio(base)

    detalhe = [
        grupo["alimentacoes"][SECAO_REGULAR]["Lanche 5h"]
        for grupo in completo["grupos"]
    ]
    consolidado = _consolidado(completo)[SECAO_REGULAR]["Lanche 5h"]
    assert consolidado == soma_colunas(detalhe, len(completo["colunas"]))

    regular = relatorio["total_geral"][SECAO_REGULAR]
    assert regular[DEZEMBRO] == 6
    assert relatorio["total_geral"][SECAO_DIETA_TIPO_A][DEZEMBRO] == 2

    linhas = _linhas(gera_relatorio_alimentacoes_servidas_xlsx(relatorio))
    total_geral = next(linha for linha in linhas if linha[0] == "TOTAL GERAL")
    assert total_geral[DEZEMBRO + 1] == 6 + 2
    assert total_geral[-1] == 6 + 2


def _lanche_emei(base, meses, **extras):
    relatorio = _relatorio(
        base,
        **{"meses[]": meses, "unidades_educacionais[]": _uuids(base["emei"])},
        **extras,
    )
    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    return relatorio, grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"]


@pytest.mark.parametrize(
    "meses, mensagem",
    [
        ([""], "Selecione ao menos um mês de referência."),
        (["2023-12"], "MM_AAAA"),
        (["13_2023"], "MM_AAAA"),
        (["12_2023", "01_2020"], "não permitido"),
    ],
)
def test_exportar_rejeita_meses_vazio_invalido_ou_nao_permitido(
    client, base, meses, mensagem
):
    client.force_login(base["usuario"])

    resposta = client.get(URL, _params(base, **{"meses[]": meses}))

    assert resposta.status_code == status.HTTP_400_BAD_REQUEST
    assert mensagem in resposta.json()["meses"][0]


def test_exportar_com_parametro_mes_antigo_retorna_400(client, base):
    client.force_login(base["usuario"])

    resposta = client.get(URL, {"mes": "12_2023", "dres[]": _uuids(base["dre"])})

    assert resposta.status_code == status.HTTP_400_BAD_REQUEST
    assert resposta.json()["meses"] == ["Selecione ao menos um mês de referência."]


def test_exportar_le_meses_repetidos_e_ordena(client, base):
    client.force_login(base["usuario"])

    with patch(DELAY) as delay:
        resposta = client.get(
            URL, _params(base, **{"meses[]": ["12_2023", "10_2023", "12_2023"]})
        )

    assert resposta.status_code == status.HTTP_200_OK
    chamada = delay.call_args.kwargs
    assert chamada["filtros"]["meses"] == ["10_2023", "12_2023"]
    assert (
        chamada["nome_arquivo"]
        == "Relatório de Alimentações Servidas - 10/2023, 12/2023.xlsx"
    )


def test_meses_duplicados_nao_duplicam_valores(base):
    relatorio, lanche = _lanche_emei(base, ["12_2023", "12_2023"])

    assert relatorio["colunas"] == ["DEZ 2023"]
    assert lanche == [10 + 5]


def test_somente_meses_selecionados_tem_coluna_e_valor(base):
    relatorio, lanche = _lanche_emei(base, ["10_2023", "12_2023"])

    assert relatorio["colunas"] == ["OUT 2023", "DEZ 2023"]
    assert relatorio["coluna_total"] == "TOTAL 2023"
    assert lanche == [3 + 4, 10 + 5]

    linhas = _linhas(gera_relatorio_alimentacoes_servidas_xlsx(relatorio))
    cabecalho = next(linha for linha in linhas if linha[0] == "Tipo de Alimentação")
    assert cabecalho == ("Tipo de Alimentação", "OUT 2023", "DEZ 2023", "TOTAL 2023")
    assert next(linha for linha in linhas if linha[0] == "Lanche 5h") == (
        "Lanche 5h",
        7,
        15,
        22,
    )
    textos = _textos(gera_relatorio_alimentacoes_servidas_xlsx(relatorio))
    assert "DRE IPIRANGA" in textos
    assert not any("DRE:" in str(texto) for texto in textos)
    assert not any("OUTUBRO" in str(texto).upper() for texto in textos)


def test_cabecalho_mostra_dres_e_periodo_sem_meses_nem_prefixo(base):
    relatorio = _relatorio(
        base,
        periodo_lancamento_de="01/12/2023",
        periodo_lancamento_ate="10/12/2023",
    )
    arquivo = gera_relatorio_alimentacoes_servidas_xlsx(relatorio)
    textos = _textos(arquivo)

    texto_filtros = "DRE IPIRANGA | PERÍODO DE LANÇAMENTO: DE 01/12/2023 ATÉ 10/12/2023"
    assert texto_filtros in textos
    aba = _abas(relatorio).worksheets[0]
    celula = next(
        celula
        for linha in aba.iter_rows()
        for celula in linha
        if celula.value == texto_filtros
    )
    assert celula.font.b is True


def test_periodo_entre_meses_diferentes_exclui_mes_nao_selecionado(base):
    relatorio, lanche = _lanche_emei(
        base,
        ["10_2023", "12_2023"],
        periodo_lancamento_de="10/10/2023",
        periodo_lancamento_ate="20/12/2023",
    )

    assert relatorio["colunas"] == ["OUT 2023", "DEZ 2023"]
    assert lanche == [4, 10 + 5]


def test_periodo_com_data_fora_dos_meses_selecionados(client, base):
    client.force_login(base["usuario"])

    resposta = client.get(
        URL,
        _params(
            base,
            periodo_lancamento_de="10/11/2023",
            periodo_lancamento_ate="20/12/2023",
            **{"meses[]": ["10_2023", "12_2023"]},
        ),
    )

    assert resposta.status_code == status.HTTP_400_BAD_REQUEST
    assert resposta.json() == {"periodo_lancamento_de": [MENSAGEM_DATA_FORA]}


def test_totais_e_consolidacoes_com_varios_meses(base):
    relatorio = _relatorio(base, **{"meses[]": ["10_2023", "12_2023"]})
    tamanho = len(relatorio["colunas"])
    consolidado = _consolidado(relatorio)

    for secao in (SECAO_REGULAR, SECAO_DIETA_TIPO_A, SECAO_DIETA_TIPO_B):
        for nome, valores in consolidado[secao].items():
            detalhe = [
                grupo["alimentacoes"][secao].get(nome, [0] * tamanho)
                for grupo in relatorio["grupos"]
            ]
            assert valores == soma_colunas(detalhe, tamanho)
        assert relatorio["total_geral"][secao] == soma_colunas(
            consolidado[secao].values(), tamanho
        )

    lanche = consolidado[SECAO_REGULAR]["Lanche 5h"]
    assert lanche == [3 + 4, 10 + 5 + 4 + 50 + 6]

    linhas = _linhas(gera_relatorio_alimentacoes_servidas_xlsx(relatorio))
    totais_por_aba = [linha[1:] for linha in linhas if linha[0] == "TOTAL GERAL"]
    esperado = soma_colunas(relatorio["total_geral"].values(), tamanho)
    assert len(totais_por_aba) == len(relatorio["grupos"])
    assert sum(total[-1] for total in totais_por_aba) == sum(esperado)


def test_selecao_entre_anos_rotula_colunas_com_o_ano(client, base):
    client.force_login(base["usuario"])

    with patch(DELAY) as delay:
        resposta = client.get(URL, _params(base, **{"meses[]": ["01_2024", "12_2023"]}))

    assert resposta.status_code == status.HTTP_200_OK
    filtros = delay.call_args.kwargs["filtros"]
    assert filtros["meses"] == ["12_2023", "01_2024"]
    assert (
        delay.call_args.kwargs["nome_arquivo"]
        == "Relatório de Alimentações Servidas - 12/2023, 01/2024.xlsx"
    )

    relatorio, lanche = _lanche_emei(base, filtros["meses"])
    assert relatorio["colunas"] == ["DEZ 2023", "JAN 2024"]
    assert relatorio["coluna_total"] == "TOTAL DO PERÍODO"
    assert lanche == [10 + 5, 11]


def test_lancamento_sem_tipo_de_alimentacao_usa_o_nome_campo(base):
    medicao = _medicao(base, base["ceu_emei"], mes="11", ano="2023")
    _valor(medicao, base["regular"], 40, dia="05", nome_campo="lanche")
    _valor(medicao, base["tipo_b"], 6, dia="05", nome_campo="refeicao")
    _valor(medicao, base["regular"], 99, dia="05", nome_campo="repeticao_refeicao")
    _valor(medicao, base["regular"], 1000, dia="05", nome_campo="frequencia")
    filtros_base = {
        "meses[]": ["11_2023"],
        "unidades_educacionais[]": _uuids(base["ceu_emei"]),
    }

    relatorio = _relatorio(base, **filtros_base)
    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    assert grupo_3["alimentacoes"][SECAO_REGULAR] == {"Lanche 5h": [40]}
    assert grupo_3["alimentacoes"][SECAO_DIETA_TIPO_B] == {"Refeição": [6]}

    filtrado = _relatorio(
        base, **filtros_base, **{"tipos_alimentacao[]": _uuids(base["refeicao"])}
    )
    grupo_3 = _grupo_do_relatorio(filtrado, GrupoUnidadeEscolar.GRUPO_3)
    assert grupo_3["alimentacoes"][SECAO_REGULAR] == {}
    assert grupo_3["alimentacoes"][SECAO_DIETA_TIPO_B] == {"Refeição": [6]}


def _abas(relatorio):
    arquivo = gera_relatorio_alimentacoes_servidas_xlsx(relatorio)
    return openpyxl.load_workbook(io.BytesIO(arquivo))


def _coluna_a(aba):
    return [aba.cell(linha, 1).value for linha in range(1, aba.max_row + 1)]


def _linha_do_texto(aba, texto, a_partir=1):
    for linha in range(a_partir, aba.max_row + 1):
        if aba.cell(linha, 1).value == texto:
            return linha
    raise AssertionError(f"{texto!r} não encontrado")


ABA_GRUPO_2 = "CEMEI"
ABA_GRUPO_3 = "EMEI, CEU EMEI"


def _total_geral(aba):
    return next(
        linha for linha in aba.iter_rows(values_only=True) if linha[0] == "TOTAL GERAL"
    )


def test_cada_aba_mostra_so_as_colunas_com_lancamento_do_grupo(base):
    relatorio = _relatorio(base, **{"meses[]": ["10_2023", "12_2023"]})
    abas = _abas(relatorio)

    assert relatorio["colunas"] == ["OUT 2023", "DEZ 2023"]
    cabecalhos = {
        aba.title: next(
            linha
            for linha in aba.iter_rows(values_only=True)
            if linha[0] == "Tipo de Alimentação"
        )
        for aba in abas.worksheets
    }
    assert cabecalhos[ABA_GRUPO_3] == (
        "Tipo de Alimentação",
        "OUT 2023",
        "DEZ 2023",
        "TOTAL 2023",
    )
    assert cabecalhos[ABA_GRUPO_2] == ("Tipo de Alimentação", "DEZ 2023", "TOTAL 2023")
    cemei = abas[ABA_GRUPO_2]
    assert _total_geral(cemei)[1:] == (6 + 2, 6 + 2)


def test_aba_unica_omite_meses_sem_lancamento(base):
    relatorio = _relatorio(
        base,
        **{
            "meses[]": ["10_2023", "12_2023"],
            "tipos_unidades[]": _uuids(base["tipo_cemei"]),
        },
    )

    linhas = _linhas(gera_relatorio_alimentacoes_servidas_xlsx(relatorio))

    cabecalho = next(linha for linha in linhas if linha[0] == "Tipo de Alimentação")
    assert cabecalho == ("Tipo de Alimentação", "DEZ 2023", "TOTAL 2023")


def test_varios_grupos_geram_uma_aba_por_grupo_nomeada_com_as_unidades(base):
    relatorio = _relatorio(base)
    abas = _abas(relatorio)

    assert abas.sheetnames == [ABA_GRUPO_2, ABA_GRUPO_3]
    for aba in abas.worksheets:
        coluna = _coluna_a(aba)
        assert coluna[0] == "Relatório de Alimentações Servidas"
        assert len(aba._images) == 1
        assert coluna.index(aba.title) < coluna.index("CONSOLIDAÇÃO - TOTAL GERAL")
        assert "CONSOLIDAÇÃO - TIPOS DE ALIMENTAÇÃO" not in coluna

    tamanho = len(relatorio["colunas"])
    global_ = soma_colunas(relatorio["total_geral"].values(), tamanho)
    por_aba = [_total_geral(aba)[1] for aba in abas.worksheets]
    assert por_aba[0] != por_aba[1]
    assert sum(por_aba) == global_[0]


def test_nome_da_aba_e_unico_e_valido(base):
    from src.medicao_inicial.services.relatorio_alimentacoes_servidas_excel import (
        _Abas,
        _nome_aba,
    )

    assert _nome_aba("EMEF: A/B?") == "EMEF  A B"
    assert len(_nome_aba("X" * 60)) == 31
    abas = _Abas.__new__(_Abas)
    abas.nomes = set()
    longo = "EMEF, EMEFM, CEU EMEF, CEU GESTAO, CIEJA"
    primeiro, segundo = abas._nome_unico(longo), abas._nome_unico(longo)
    assert primeiro != segundo
    assert len(primeiro) <= 31 and len(segundo) <= 31
    assert segundo.endswith("(2)")


def test_um_grupo_fica_em_uma_aba_com_a_consolidacao(base):
    relatorio = _relatorio(base, **{"unidades_educacionais[]": _uuids(base["emei"])})
    abas = _abas(relatorio)

    assert abas.sheetnames == ["Alimentações Servidas"]
    coluna = _coluna_a(abas.active)
    assert "EMEI" in coluna
    assert "CONSOLIDAÇÃO - TOTAL GERAL" in coluna
    assert coluna.index("EMEI") < coluna.index("CONSOLIDAÇÃO - TOTAL GERAL")


def test_grupo_sem_linha_de_total_por_secao_e_so_o_total_geral_na_consolidacao(base):
    coluna = _coluna_a(_abas(_relatorio(base))[ABA_GRUPO_3])
    inicio = coluna.index("CONSOLIDAÇÃO - TOTAL GERAL")
    detalhe, consolidacao = coluna[:inicio], coluna[inicio:]

    assert "TOTAL" not in detalhe
    assert detalhe.count("Tipo de Alimentação") == 1
    assert "Tipo de Alimentação" not in consolidacao
    assert consolidacao.count("TOTAL") == 0
    assert consolidacao.count("TOTAL GERAL") == 1


def test_secoes_coladas_na_ordem_dentro_do_grupo(base):
    aba = _abas(_relatorio(base))[ABA_GRUPO_3]
    coluna = _coluna_a(aba)

    titulo = coluna.index("EMEI, CEU EMEI")
    assert coluna[titulo + 1] == "Tipo de Alimentação"
    posicoes = [
        coluna.index(secao)
        for secao in (SECAO_REGULAR, SECAO_DIETA_TIPO_A, SECAO_DIETA_TIPO_B)
    ]
    assert posicoes == sorted(posicoes)
    assert coluna[titulo + 2 : posicoes[-1] + 1].count(None) == 0


def test_cores_distintas_por_secao_e_titulo_do_grupo_cinza(base):
    aba = _abas(_relatorio(base))[ABA_GRUPO_3]

    cores = {
        secao: aba.cell(_linha_do_texto(aba, secao), 1).fill.start_color.rgb[-6:]
        for secao in (SECAO_REGULAR, SECAO_DIETA_TIPO_A, SECAO_DIETA_TIPO_B)
    }
    assert len(set(cores.values())) == 3
    assert cores[SECAO_DIETA_TIPO_A] == "B1DBB7"
    assert cores[SECAO_DIETA_TIPO_B] == "67A37E"

    titulo = aba.cell(_linha_do_texto(aba, "EMEI, CEU EMEI"), 1)
    assert titulo.fill.start_color.rgb[-6:] == "D9D9D9"
    assert titulo.font.b is True
    assert titulo.font.color.rgb[-6:] == "000000"


def test_logo_fica_centralizada_na_vertical_do_titulo(base):
    for aba in _abas(_relatorio(base)).worksheets:
        assert aba._images[0].anchor._from.rowOff > 5 * 9525
        assert aba.row_dimensions[1].height == 70


def test_mes_regular_nao_soma_o_recreio_e_o_recreio_nao_soma_o_regular(base):
    _, so_regular = _lanche_emei(base, ["01_2024"])
    assert so_regular == [11]

    so_recreio = _relatorio(
        base,
        **{
            "meses[]": [],
            "recreios_nas_ferias[]": _uuids(base["recreio"]),
            "unidades_educacionais[]": _uuids(base["emei"]),
        },
    )
    grupo_3 = _grupo_do_relatorio(so_recreio, GrupoUnidadeEscolar.GRUPO_3)
    assert so_recreio["colunas"] == ["RECREIO NAS FÉRIAS - JAN 2024"]
    assert grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"] == [9]
    assert so_recreio["coluna_total"] == "TOTAL 2024"


def test_periodo_entre_mes_e_recreio_nao_inclui_mes_nao_selecionado(base):
    relatorio = _com_recreio(
        base,
        ["12_2023"],
        periodo_lancamento_de="20/12/2023",
        periodo_lancamento_ate="05/01/2024",
    )
    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)

    assert relatorio["colunas"] == ["DEZ 2023", "RECREIO NAS FÉRIAS - JAN 2024"]
    assert grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"] == [5, 9]


def test_recreios_do_mesmo_mes_recebem_o_titulo_na_coluna():
    recreios = [
        {"id": 1, "titulo": "Recreio A", "meses": [(1, 2024)]},
        {"id": 2, "titulo": "Recreio B", "meses": [(1, 2024)]},
    ]

    colunas = monta_colunas([(1, 2024)], recreios)

    assert [coluna["rotulo"] for coluna in colunas] == [
        "REGULAR - JAN 2024",
        "RECREIO NAS FÉRIAS - JAN 2024 (Recreio A)",
        "RECREIO NAS FÉRIAS - JAN 2024 (Recreio B)",
    ]
    assert [coluna["chave"] for coluna in colunas] == [
        (2024, 1, None),
        (2024, 1, 1),
        (2024, 1, 2),
    ]


def test_quantidade_de_consultas_nao_cresce_com_meses_e_recreios(base):
    def consultas(meses):
        with CaptureQueriesContext(connection) as contexto:
            _com_recreio(base, meses)
        return len(contexto)

    consultas(["12_2023"])  # aquece os caches do Django, que enviesam a contagem
    assert consultas(["12_2023", "01_2024"]) == consultas(
        ["10_2023", "11_2023", "12_2023", "01_2024"]
    )


def test_exportar_com_mes_e_recreio_registra_nome_e_cabecalho(client, base):
    client.force_login(base["usuario"])

    resposta = client.get(
        URL, _params(base, **{"recreios_nas_ferias[]": _uuids(base["recreio"])})
    )

    assert resposta.status_code == status.HTTP_200_OK
    central = CentralDeDownload.objects.get(usuario=base["usuario"])
    assert central.identificador == (
        "Relatório de Alimentações Servidas - 12/2023, Recreio (01/2024).xlsx"
    )
    with central.arquivo.open("rb") as arquivo:
        textos = _textos(arquivo.read())
    assert "DRE IPIRANGA" in textos
    assert not any("| RECREIO" in str(texto) for texto in textos)
    assert "RECREIO NAS FÉRIAS - JAN 2024" in textos
    assert "DEZ 2023" in textos


def test_exportar_aceita_so_recreio_e_valida_os_recreios(client, base):
    client.force_login(base["usuario"])
    dres = {"dres[]": _uuids(base["dre"])}

    with patch(DELAY) as delay:
        resposta = client.get(
            URL, {"recreios_nas_ferias[]": _uuids(base["recreio"]), **dres}
        )

    assert resposta.status_code == status.HTTP_200_OK
    chamada = delay.call_args.kwargs
    assert chamada["filtros"]["meses"] == []
    assert chamada["filtros"]["recreios_nas_ferias"] == _uuids(base["recreio"])
    assert chamada["nome_arquivo"] == (
        "Relatório de Alimentações Servidas - Recreio (01/2024).xlsx"
    )

    inexistente = {"recreios_nas_ferias[]": [str(uuid4())]}
    so_recreio = client.get(URL, {**inexistente, **dres})
    assert so_recreio.status_code == status.HTTP_400_BAD_REQUEST
    assert "recreios_nas_ferias" in so_recreio.json()
    com_mes = client.get(URL, _params(base, **inexistente))
    assert com_mes.status_code == status.HTTP_400_BAD_REQUEST
    assert "recreios_nas_ferias" in com_mes.json()
    invalido = client.get(URL, {"recreios_nas_ferias[]": ["abc"], **dres})
    assert invalido.status_code == status.HTTP_400_BAD_REQUEST
    assert "recreios_nas_ferias" in invalido.json()


def test_nome_do_arquivo_cabe_no_identificador_da_central():
    meses = [f"{mes:02d}_{ano}" for ano in range(2015, 2026) for mes in range(1, 13)]

    nome = obtem_nome_arquivo_relatorio_alimentacoes_servidas(meses)

    assert len(nome) <= 200
    assert nome.startswith("Relatório de Alimentações Servidas - 01/2015")
    assert nome.endswith("….xlsx")


def _lancamentos_grupo_1(base):
    cei = _escola(base, base["dre"], base["lote"], base["tipo_cei"], base["sub"])
    medicao = _medicao(base, cei)
    _valor(medicao, base["regular"], 3, alimentacao=base["lanche"])


def test_varios_grupos_de_unidades_so_restringem_alimentacao_fora_do_grupo_1(base):
    _lancamentos_grupo_1(base)

    relatorio = _relatorio(
        base,
        **{
            "tipos_unidades[]": _uuids(base["tipo_cei"], base["tipo_emei"]),
            "tipos_alimentacao[]": _uuids(base["refeicao"]),
        },
    )

    assert {grupo["nome"] for grupo in relatorio["grupos"]} == {
        GrupoUnidadeEscolar.GRUPO_1,
        GrupoUnidadeEscolar.GRUPO_3,
    }
    grupo_1 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_1)
    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    assert grupo_1["alimentacoes"][SECAO_REGULAR] == {"Lanche 5h": [3]}
    assert grupo_3["alimentacoes"][SECAO_REGULAR] == {"Refeição": [7]}
    assert grupo_3["tipos_unidades"] == ["EMEI"]


CHAVES_REGULARES = {
    "Refeição": "total_refeicao",
    "Refeição EJA": "total_refeicao_eja",
    "Sobremesa": "total_sobremesa",
    "Lanche 4h": "total_lanche_4h",
    "Lanche 5h": "total_lanche",
    "Lanche Emergencial": "total_lanche_emergencial",
    "Kit Lanche Passeio": "total_kit_lanche",
}
CHAVES_DIETA = {
    "Refeição": "refeicao",
    "Refeição EJA": "refeicao_eja",
    "Lanche 5h": "lanche",
    "Lanche 4h": "lanche_4h",
}


def _esperado_do_endpoint(base, nome_grupo, mes, ano):
    grupo = GrupoUnidadeEscolar.objects.get(nome=nome_grupo)
    return calcula_totais_consumo_por_grupo(
        base["lote"], grupo, mes, ano, "tipo_alimentacao"
    )


def _confere_com_o_endpoint(grupo, esperado, coluna=DEZEMBRO):
    """Cada linha do grupo no Excel é a chave correspondente do endpoint."""
    secoes = {
        SECAO_REGULAR: (CategoriaMedicao.ALIMENTACAO, CHAVES_REGULARES),
        SECAO_DIETA_TIPO_A: (CategoriaMedicao.DIETA_ESPECIAL_TIPO_A, CHAVES_DIETA),
        SECAO_DIETA_TIPO_B: (CategoriaMedicao.DIETA_ESPECIAL_TIPO_B, CHAVES_DIETA),
    }
    for secao, (categoria, chaves) in secoes.items():
        linhas = grupo["alimentacoes"][secao]
        do_endpoint = esperado.get(categoria, {})
        assert set(linhas) == {r for r, k in chaves.items() if do_endpoint.get(k)}
        for rotulo, valores in linhas.items():
            assert valores[coluna] == do_endpoint[chaves[rotulo]], (secao, rotulo)


def test_valores_batem_com_o_endpoint_de_totais_de_atendimento_e_consumo(base):
    solicitacoes = baker.make(
        "CategoriaMedicao", nome=CategoriaMedicao.SOLICITACOES_DE_ALIMENTACAO
    )
    medicao = _medicao(base, base["emei"], mes="05", ano="2026")
    _valor(medicao, base["regular"], 12, alimentacao=base["lanche"])
    _valor(medicao, base["regular"], 9, alimentacao=base["refeicao"])
    _valor(medicao, base["regular"], 3, nome_campo="2_refeicao_1_oferta")
    _valor(medicao, base["regular"], 5, nome_campo="repeticao_refeicao")
    _valor(medicao, base["regular"], 6, nome_campo="lanche_4h")
    _valor(medicao, base["regular"], 2, nome_campo="2_lanche_5h")
    _valor(medicao, base["regular"], 4, nome_campo="sobremesa")
    _valor(medicao, base["regular"], 100, nome_campo="matriculados")
    _valor(medicao, solicitacoes, 7, nome_campo="kit_lanche")
    _valor(medicao, solicitacoes, 3, nome_campo="lanche_emergencial")
    _valor(medicao, base["tipo_a"], 2, alimentacao=base["lanche"])
    _valor(medicao, base["tipo_a"], 8, nome_campo="lanche_4h")
    _valor(medicao, base["tipo_b"], 1, alimentacao=base["refeicao"])
    parametros = {
        "meses[]": ["05_2026"],
        "lotes[]": _uuids(base["lote"]),
        "tipos_unidades[]": _uuids(base["tipo_emei"], base["tipo_ceu_emei"]),
    }

    relatorio = _relatorio(base, **parametros)

    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    esperado = _esperado_do_endpoint(base, GrupoUnidadeEscolar.GRUPO_3, "05", "2026")
    assert esperado[CategoriaMedicao.ALIMENTACAO]["total_refeicao"] == 9 + 3
    _confere_com_o_endpoint(grupo_3, esperado, coluna=0)
    regulares = grupo_3["alimentacoes"][SECAO_REGULAR]
    assert regulares["Lanche 5h"] == [12 + 2]
    assert regulares["Kit Lanche Passeio"] == [7]
    assert list(regulares) == [
        "Refeição",
        "Sobremesa",
        "Lanche 4h",
        "Lanche 5h",
        "Lanche Emergencial",
        "Kit Lanche Passeio",
    ]


def test_dezembro_do_grupo_3_bate_com_o_endpoint_e_consolidado_soma_regular_e_dietas(
    base,
):
    relatorio = _relatorio(
        base,
        **{
            "lotes[]": _uuids(base["lote"]),
            "tipos_unidades[]": _uuids(base["tipo_emei"], base["tipo_ceu_emei"]),
        },
    )

    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    _confere_com_o_endpoint(
        grupo_3, _esperado_do_endpoint(base, GrupoUnidadeEscolar.GRUPO_3, "12", "2023")
    )
    consolidado = _consolidado(relatorio)
    assert consolidado[SECAO_REGULAR]["Lanche 5h"][DEZEMBRO] == 10 + 5 + 4
    assert grupo_3["alimentacoes"][SECAO_DIETA_TIPO_A]["Lanche 5h"] == [2]


def test_recreio_e_regular_do_mes_somados_batem_com_o_endpoint(base):
    relatorio = _com_recreio(base, ["01_2024"], **{"lotes[]": _uuids(base["lote"])})
    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    esperado = _esperado_do_endpoint(base, GrupoUnidadeEscolar.GRUPO_3, "01", "2024")

    regular, recreio = grupo_3["alimentacoes"][SECAO_REGULAR]["Lanche 5h"]

    assert regular + recreio == esperado[CategoriaMedicao.ALIMENTACAO]["total_lanche"]


def _grupo_4(base):
    tipo = baker.make("TipoUnidadeEscolar", iniciais="EMEF")
    _grupo(GrupoUnidadeEscolar.GRUPO_4, tipo)
    escola = _escola(base, base["dre"], base["lote"], tipo, base["sub"])
    noite = baker.make("PeriodoEscolar", nome="NOITE")
    return tipo, escola, noite


def _lancamentos_emef(base, escola, noite):
    manha = _medicao(base, escola, mes="05", ano="2026")
    noturno = baker.make(
        "Medicao",
        solicitacao_medicao_inicial=manha.solicitacao_medicao_inicial,
        periodo_escolar=noite,
    )
    regular = base["regular"]
    for dia, matriculados in (("01", 10), ("02", 10)):
        _valor(manha, regular, matriculados, dia=dia, nome_campo="matriculados")
    _valor(manha, regular, 8, dia="01", nome_campo="refeicao")
    _valor(manha, regular, 5, dia="01", nome_campo="repeticao_refeicao")
    _valor(manha, regular, 4, dia="02", nome_campo="refeicao")
    _valor(manha, regular, 3, dia="02", nome_campo="2_refeicao_1_oferta")
    _valor(manha, regular, 1, dia="02", nome_campo="repeticao_2_refeicao")
    _valor(manha, regular, 12, dia="01", nome_campo="sobremesa")
    _valor(manha, regular, 5, dia="01", nome_campo="lanche")
    _valor(manha, base["tipo_b"], 2, dia="01", nome_campo="refeicao")
    _valor(noturno, regular, 20, dia="01", nome_campo="matriculados")
    _valor(noturno, regular, 6, dia="01", nome_campo="refeicao")
    _valor(noturno, base["tipo_b"], 4, dia="01", nome_campo="refeicao")


def test_emef_usa_o_total_de_pagamento_limitado_aos_matriculados(base):
    tipo, escola, noite = _grupo_4(base)
    _lancamentos_emef(base, escola, noite)
    parametros = {
        "meses[]": ["05_2026"],
        "lotes[]": _uuids(base["lote"]),
        "tipos_unidades[]": _uuids(tipo),
    }

    relatorio = _relatorio(base, **parametros)

    grupo_4 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_4)
    regulares = grupo_4["alimentacoes"][SECAO_REGULAR]
    assert regulares["Refeição"] == [min(8 + 5, 10) + min(4, 10) + min(3 + 1, 10)]
    assert regulares["Refeição EJA"] == [6]
    assert regulares["Sobremesa"] == [10]
    assert regulares["Lanche 5h"] == [5]
    assert grupo_4["alimentacoes"][SECAO_DIETA_TIPO_B] == {
        "Refeição": [2],
        "Refeição EJA": [4],
    }
    esperado = _esperado_do_endpoint(base, GrupoUnidadeEscolar.GRUPO_4, "05", "2026")
    _confere_com_o_endpoint(grupo_4, esperado, coluna=0)


def test_emef_no_recreio_compara_com_participantes(base):
    tipo, escola, _ = _grupo_4(base)
    recreio = baker.make(
        "RecreioNasFerias",
        titulo="Recreio EMEF",
        data_inicio=datetime.date(2024, 1, 2),
        data_fim=datetime.date(2024, 1, 31),
    )
    medicao = _medicao(base, escola, mes="01", ano="2024", recreio=recreio)
    _valor(medicao, base["regular"], 7, nome_campo="participantes")
    _valor(medicao, base["regular"], 9, nome_campo="refeicao")
    _valor(medicao, base["regular"], 99, nome_campo="matriculados")

    relatorio = _relatorio(
        base,
        **{
            "meses[]": [],
            "recreios_nas_ferias[]": _uuids(recreio),
            "tipos_unidades[]": _uuids(tipo),
        },
    )

    grupo_4 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_4)
    assert grupo_4["alimentacoes"][SECAO_REGULAR]["Refeição"] == [7]


def test_tipos_alimentacao_filtra_pelo_campo_consolidado(base):
    relatorio = _relatorio(
        base,
        **{
            "unidades_educacionais[]": _uuids(base["emei"]),
            "tipos_alimentacao[]": _uuids(base["lanche"]),
        },
    )

    grupo_3 = _grupo_do_relatorio(relatorio, GrupoUnidadeEscolar.GRUPO_3)
    assert list(grupo_3["alimentacoes"][SECAO_REGULAR]) == ["Lanche 5h"]
    assert list(grupo_3["alimentacoes"][SECAO_DIETA_TIPO_A]) == ["Lanche 5h"]


def test_exportar_aceita_tipos_de_unidade_de_grupos_diferentes(client, base):
    client.force_login(base["usuario"])
    tipos = _uuids(base["tipo_cemei"], base["tipo_emei"])

    with patch(DELAY) as delay:
        resposta = client.get(URL, _params(base, **{"tipos_unidades[]": tipos}))

    assert resposta.status_code == status.HTTP_200_OK
    assert set(delay.call_args.kwargs["filtros"]["tipos_unidades"]) == set(tipos)
    relatorio = monta_relatorio_alimentacoes_servidas(
        delay.call_args.kwargs["filtros"], base["usuario"]
    )
    assert {grupo["nome"] for grupo in relatorio["grupos"]} == {
        GrupoUnidadeEscolar.GRUPO_2,
        GrupoUnidadeEscolar.GRUPO_3,
    }
