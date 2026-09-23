from django.core.management.base import BaseCommand
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from src.dados_comuns.constants import NOME_PROTOCOLO_PLANILHA, StringsNomesAbasXLSX
from src.dieta_especial.protocolo_padrao.models import (
    ProtocoloPadraoDietaEspecial,
)
from src.dieta_especial.solicitacao_dieta_especial.models import (
    SolicitacaoDietaEspecial,
)
from src.terceirizada.models import Edital


class Command(BaseCommand):
    help = """
    Processar planilhas de devolutiva das POs para Dietas Especiais com protocolos inválidos
    """

    def extrair_dados_planilha(self):
        active_sheet = load_workbook("invalidos.xlsx").active
        linhas = list(active_sheet.rows)
        lista_objetos = []
        for linha in enumerate(linhas[2:]):
            objeto = {}
            for c, coluna in enumerate(linha[1]):
                objeto[linhas[0][c].value] = coluna.value
            lista_objetos.append(objeto)
        return lista_objetos

    def formatar_tamanho_celulas(self, ws):
        for column in ["A", "B", "C", "D", "E", "F"]:
            ws.column_dimensions[column].width = 50

    def exportar_planilha(self, lista):
        wb = Workbook()
        ws = wb.active
        self.formatar_tamanho_celulas(ws)
        ws.title = StringsNomesAbasXLSX.DIETAS_NAO_RELACIONADAS.value
        cabecalho = [
            "uuid",
            NOME_PROTOCOLO_PLANILHA,
            "escola",
            "aluno",
            "lote",
            "erro",
        ]
        for ind, title in enumerate(cabecalho, 1):
            celula = ws.cell(row=1, column=ind)
            celula.value = str(title)
            celula.font = Font(size="13", bold=True)

        for ind, dict_solicitacao in enumerate(lista, 2):
            ws.cell(row=ind, column=1, value=str(dict_solicitacao["uuid"]))
            ws.cell(row=ind, column=2, value=dict_solicitacao[NOME_PROTOCOLO_PLANILHA])
            ws.cell(row=ind, column=3, value=dict_solicitacao["escola"])
            ws.cell(row=ind, column=4, value=dict_solicitacao["aluno"])
            ws.cell(row=ind, column=5, value=dict_solicitacao["lote"])
            ws.cell(row=ind, column=6, value=dict_solicitacao["erro"])

        wb.save("relacao-planilha-invalidos-protocolos-com-erro.xlsx")

    def check_deletar_solicitacoes(self, nome_planilha, nome_db):
        if nome_planilha is None and nome_db is None:
            return True
        else:
            return False

    def handle(self, *args, **options):
        dados_planilha = self.extrair_dados_planilha()
        solicitacoes_nao_relacionadas = []
        for dado in dados_planilha:
            self._processar_dado(dado, solicitacoes_nao_relacionadas)
        self.exportar_planilha(solicitacoes_nao_relacionadas)

    def _processar_dado(self, dado, solicitacoes_nao_relacionadas):
        solicitacao = SolicitacaoDietaEspecial.objects.filter(uuid=dado["UUID"]).first()
        if not solicitacao:
            return
        nome_planilha = dado["ALTERAR NOME DO PROTOCOLO NO SISTEMA  PARA:"]
        nome_db = dado["CRIADO PROTOCOLO PADRÃO COM NOME:"]
        if self.check_deletar_solicitacoes(nome_planilha, nome_db):
            solicitacao.delete()
            return
        nome_protocolo = nome_planilha if nome_planilha else nome_db
        protocolos = self._buscar_protocolos(solicitacao, nome_protocolo)
        if protocolos:
            solicitacao.protocolo_padrao = protocolos.first()
            solicitacao.save()
            return
        solicitacoes_nao_relacionadas.append(
            self._montar_objeto(dado, solicitacao, nome_protocolo)
        )

    def _buscar_protocolos(self, solicitacao, nome_protocolo):
        editais_uuids = solicitacao.escola.lote.contratos_do_lote.values_list(
            "edital__uuid", flat=True
        )
        editais = Edital.objects.filter(uuid__in=editais_uuids)
        protocolos_uuids = editais.values_list(
            "protocolos_padroes_dieta_especial__uuid"
        )
        return ProtocoloPadraoDietaEspecial.objects.filter(
            uuid__in=protocolos_uuids, nome_protocolo=nome_protocolo
        )

    def _montar_objeto(self, dado, solicitacao, nome_protocolo):
        return {
            "uuid": dado["UUID"],
            NOME_PROTOCOLO_PLANILHA: nome_protocolo,
            "escola": solicitacao.escola.nome,
            "aluno": solicitacao.aluno.nome,
            "lote": solicitacao.escola.lote.nome,
            "erro": "Protocolo não encontrado para o Edital e Lote da Escola ",
        }
