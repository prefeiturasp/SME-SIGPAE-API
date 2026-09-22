import os
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from src.dados_comuns.constants import (
    FORMATO_DATA_BRASILEIRO,
    StringsCabecalhoXLSXGuiaDaRemessa,
)

from ..helpers import (
    retorna_motivo_insucesso,
    retorna_ocorrencias_alimento,
    retorna_status_alimento,
    retorna_status_guia_remessa,
)


class RequisicoesExcelService(object):
    DEFAULT_BORDER = Border(
        right=Side(border_style="thin", color="24292E"),
        left=Side(border_style="thin", color="24292E"),
        top=Side(border_style="thin", color="24292E"),
        bottom=Side(border_style="thin", color="24292E"),
    )

    @classmethod
    def aplicar_tamanho_calculado_nas_celulas(cls, ws):
        # Ajuste automatico do tamanho das colunas
        for colunas in ws.columns:
            unmerged_cells = list(
                filter(
                    lambda cell_to_check: cell_to_check.coordinate
                    not in ws.merged_cells,
                    colunas,
                )
            )
            length = max(len(str(cell.value)) for cell in unmerged_cells)
            ws.column_dimensions[unmerged_cells[0].column_letter].width = length * 1.2

    @classmethod
    def aplicar_estilo_padrao(cls, ws, count_data, count_fields):
        for linha in range(1, (count_data + 2)):
            for coluna in range(1, (count_fields + 1)):
                celula = ws.cell(row=linha, column=coluna)
                celula.border = cls.DEFAULT_BORDER
                if linha == 1:
                    celula.fill = PatternFill(fill_type="solid", fgColor="198459")

        cls.aplicar_tamanho_calculado_nas_celulas(ws)

    @classmethod
    def aplicar_estilo_visao_distribuidor(cls, ws, count_data, count_fields):
        for linha in range(2, (count_data + 3)):
            for coluna in range(1, (count_fields + 1)):
                celula = ws.cell(row=linha, column=coluna)
                celula.border = cls.DEFAULT_BORDER
                if linha == 2:
                    celula.fill = PatternFill(fill_type="solid", fgColor="198459")

        cls.aplicar_tamanho_calculado_nas_celulas(ws)

    @classmethod
    def gera_arquivo(cls, wb):
        with NamedTemporaryFile(delete=False) as tmp:
            wb.save(tmp.name)
            tmp.seek(0)
            arquivo = tmp.read()
            tmp.close()
            os.unlink(tmp.name)
            return arquivo

    @classmethod  # noqa C901
    def exportar_visao_distribuidor(cls, requisicoes, is_async=False):
        cabecalho = [
            StringsCabecalhoXLSXGuiaDaRemessa.NO_DA_REQUISICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DA_REQUISICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_DE_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.ALIMENTO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_EOL_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.ENDERECO_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.BAIRRO_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CEP_DA_UE_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.TELEFONE_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_DA_GUIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CAPACIDADE_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CAPACIDADE_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_SUPRI_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.AGRUPAMENTO.value,
        ]

        count_fields = len(cabecalho)
        count_data = requisicoes.count()

        wb = Workbook()
        ws = wb.active
        ws.title = "Visão Analítica Abastecimento"

        for ind, title in enumerate(cabecalho, 1):
            celula = ws.cell(row=2, column=ind)
            celula.value = title
            celula.font = Font(size="13", bold=True, color="00FFFFFF")

        for ind, requisicao in enumerate(requisicoes, 3):
            ws.cell(row=ind, column=1, value=int(requisicao["numero_solicitacao"]))
            ws.cell(row=ind, column=2, value=requisicao["status_requisicao"])
            ws.cell(
                row=ind,
                column=3,
                value=requisicao["guias__data_entrega"].strftime(
                    FORMATO_DATA_BRASILEIRO
                ),
            )
            ws.cell(
                row=ind, column=4, value=requisicao["guias__alimentos__nome_alimento"]
            )
            ws.cell(row=ind, column=5, value=requisicao["guias__nome_unidade"])
            ws.cell(row=ind, column=6, value=requisicao["guias__escola__codigo_eol"])
            ws.cell(
                row=ind,
                column=7,
                value=f'{requisicao["guias__endereco_unidade"]} '
                f'{requisicao["guias__numero_unidade"]}',
            )
            ws.cell(row=ind, column=8, value=requisicao["guias__bairro_unidade"])
            ws.cell(row=ind, column=9, value=requisicao["guias__cep_unidade"])
            ws.cell(row=ind, column=10, value=requisicao["guias__telefone_unidade"])
            ws.cell(row=ind, column=11, value=int(requisicao["guias__numero_guia"]))
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(
                    row=ind,
                    column=12,
                    value=f'{requisicao["guias__alimentos__embalagens__qtd_volume"]} '
                    f'{requisicao["guias__alimentos__embalagens__descricao_embalagem"]}',
                )
                ws.cell(
                    row=ind,
                    column=13,
                    value=f'{requisicao["guias__alimentos__embalagens__capacidade_embalagem"]} '
                    f'{requisicao["guias__alimentos__embalagens__unidade_medida"]}',
                )
            if (
                requisicao["guias__alimentos__embalagens__tipo_embalagem"]
                == "FRACIONADA"
            ):
                ws.cell(
                    row=ind,
                    column=14,
                    value=f'{requisicao["guias__alimentos__embalagens__qtd_volume"]} '
                    f'{requisicao["guias__alimentos__embalagens__descricao_embalagem"]}',
                )
                ws.cell(
                    row=ind,
                    column=15,
                    value=f'{requisicao["guias__alimentos__embalagens__capacidade_embalagem"]} '
                    f'{requisicao["guias__alimentos__embalagens__unidade_medida"]}',
                )
            ws.cell(
                row=ind,
                column=16,
                value=requisicao["guias__alimentos__codigo_suprimento"],
            )
            ws.cell(
                row=ind,
                column=17,
                value=requisicao["guias__escola__subprefeitura__agrupamento"],
            )

        cls.aplicar_estilo_visao_distribuidor(ws, count_data, count_fields)
        arquivo = cls.gera_arquivo(wb)
        filename = "visao-consolidada.xlsx"

        return arquivo if is_async else {"arquivo": arquivo, "filename": filename}

    @classmethod
    def exportar_visao_dilog(cls, requisicoes, is_async=False):
        cabecalho = [
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_DISTRIBUIDOR.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_DA_REQUISICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DA_REQUISICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_TOTAL_DE_GUIAS.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_DA_GUIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_DE_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_EOL_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_CODAE_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.ENDERECO_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.BAIRRO_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CEP_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CIDADE_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.ESTADO_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CONTATO_DE_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.TELEFONE_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_ALIMENTO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_SUPRI.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_PAPA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DESCRICAO_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CAPACIDADE_DA_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.UNIDADE_DE_MEDIDA_DA_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_DE_VOLUMES_DA_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DESCRICAO_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CAPACIDADE_DA_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.UNIDADE_DE_MEDIDA_DA_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_DE_VOLUMES_DA_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DA_GUIA.value,
        ]

        count_fields = len(cabecalho)
        count_data = requisicoes.count()

        wb = Workbook()
        ws = wb.active
        ws.title = "Visão Analítica Abastecimento"

        for ind, title in enumerate(cabecalho, 1):
            celula = ws.cell(row=1, column=ind)
            celula.value = title
            celula.font = Font(size="13", bold=True, color="00FFFFFF")

        for ind, requisicao in enumerate(requisicoes, 2):
            ws.cell(row=ind, column=1, value=requisicao["distribuidor__nome_fantasia"])
            ws.cell(row=ind, column=2, value=requisicao["numero_solicitacao"])
            ws.cell(row=ind, column=3, value=requisicao["status_requisicao"])
            ws.cell(row=ind, column=4, value=requisicao["quantidade_total_guias"])
            ws.cell(row=ind, column=5, value=requisicao["guias__numero_guia"])
            ws.cell(
                row=ind,
                column=6,
                value=requisicao["guias__data_entrega"].strftime(
                    FORMATO_DATA_BRASILEIRO
                ),
            )
            ws.cell(row=ind, column=7, value=requisicao["codigo_eol_unidade"])
            ws.cell(row=ind, column=8, value=requisicao["guias__codigo_unidade"])
            ws.cell(row=ind, column=9, value=requisicao["guias__nome_unidade"])
            ws.cell(row=ind, column=10, value=requisicao["guias__endereco_unidade"])
            ws.cell(row=ind, column=11, value=(requisicao["guias__numero_unidade"]))
            ws.cell(row=ind, column=12, value=requisicao["guias__bairro_unidade"])
            ws.cell(row=ind, column=13, value=requisicao["guias__cep_unidade"])
            ws.cell(row=ind, column=14, value=requisicao["guias__cidade_unidade"])
            ws.cell(row=ind, column=15, value=requisicao["guias__estado_unidade"])
            ws.cell(row=ind, column=16, value=requisicao["guias__contato_unidade"])
            ws.cell(row=ind, column=17, value=requisicao["guias__telefone_unidade"])
            ws.cell(
                row=ind, column=18, value=requisicao["guias__alimentos__nome_alimento"]
            )
            ws.cell(
                row=ind,
                column=19,
                value=requisicao["guias__alimentos__codigo_suprimento"],
            )
            ws.cell(
                row=ind, column=20, value=requisicao["guias__alimentos__codigo_papa"]
            )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(
                    row=ind,
                    column=21,
                    value=requisicao[
                        "guias__alimentos__embalagens__descricao_embalagem"
                    ],
                )
                ws.cell(
                    row=ind,
                    column=22,
                    value=requisicao[
                        "guias__alimentos__embalagens__capacidade_embalagem"
                    ],
                )
                ws.cell(
                    row=ind,
                    column=23,
                    value=requisicao["guias__alimentos__embalagens__unidade_medida"],
                )
                ws.cell(
                    row=ind,
                    column=24,
                    value=requisicao["guias__alimentos__embalagens__qtd_volume"],
                )
            else:
                ws.cell(
                    row=ind,
                    column=25,
                    value=requisicao[
                        "guias__alimentos__embalagens__descricao_embalagem"
                    ],
                )
                ws.cell(
                    row=ind,
                    column=26,
                    value=requisicao[
                        "guias__alimentos__embalagens__capacidade_embalagem"
                    ],
                )
                ws.cell(
                    row=ind,
                    column=27,
                    value=requisicao["guias__alimentos__embalagens__unidade_medida"],
                )
                ws.cell(
                    row=ind,
                    column=28,
                    value=requisicao["guias__alimentos__embalagens__qtd_volume"],
                )
            ws.cell(row=ind, column=29, value=requisicao["guias__status"])

        cls.aplicar_estilo_padrao(ws, count_data, count_fields)
        arquivo = cls.gera_arquivo(wb)

        return arquivo if is_async else {"arquivo": arquivo}

    @classmethod
    def cria_aba_insucesso(cls, ws, requisicoes, perfil):  # noqa C901
        offset = 0
        cabecalho = [
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_DA_REQUISICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_TOTAL_DE_GUIAS.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_DA_GUIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_DE_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_MOTORISTA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.PLACA_DO_VEICULO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_DE_REGISTRO_DO_INSUCESSO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.HORA_DE_REGISTRO_DO_INSUCESSO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_EOL.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_CODAE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.ENDERECO_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.BAIRRO_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CEP_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CIDADE_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.ESTADO_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CONTATO_DE_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.TELEFONE_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_ALIMENTO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_SUPRI.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_PAPA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DESCRICAO_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CAPACIDADE_DA_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.UNIDADE_DE_MEDIDA_DA_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_PREVISTA_VOLUMES_DA_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DESCRICAO_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CAPACIDADE_DA_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.UNIDADE_DE_MEDIDA_DA_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_PREVISTA_VOLUMES_DA_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.HORA_DA_TENTATIVA_DE_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.MOTIVO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.JUSTIFICATIVA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DA_GUIA_DE_REMESSA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_COMPLETO_DO_CONFERENTE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DOCUMENTO_DO_CONFERENTE.value,
        ]

        if perfil == "DILOG":
            cabecalho.insert(0, "Nome do Distribuidor")
            offset = 1

        count_fields = len(cabecalho)
        count_data = requisicoes.count()

        ws.title = "Relatório de Insucesso"
        for ind, title in enumerate(cabecalho, 1):
            celula = ws.cell(row=1, column=ind)
            celula.value = title
            celula.font = Font(size="13", bold=True, color="00FFFFFF")

        for ind, requisicao in enumerate(requisicoes, 2):
            if perfil == "DILOG":
                ws.cell(
                    row=ind,
                    column=offset,
                    value=requisicao["distribuidor__nome_fantasia"],
                )
            ws.cell(row=ind, column=offset + 1, value=requisicao["numero_solicitacao"])
            ws.cell(
                row=ind, column=offset + 2, value=requisicao["quantidade_total_guias"]
            )
            ws.cell(row=ind, column=offset + 3, value=requisicao["guias__numero_guia"])
            ws.cell(row=ind, column=offset + 4, value=requisicao["guias__data_entrega"])
            if requisicao["guias__insucessos__criado_em"] is not None:
                ws.cell(
                    row=ind,
                    column=offset + 5,
                    value=requisicao["guias__insucessos__nome_motorista"],
                )
                ws.cell(
                    row=ind,
                    column=offset + 6,
                    value=requisicao["guias__insucessos__placa_veiculo"],
                )
                ws.cell(
                    row=ind,
                    column=offset + 7,
                    value=requisicao["guias__insucessos__criado_em"].strftime(
                        FORMATO_DATA_BRASILEIRO
                    ),
                )
                ws.cell(
                    row=ind,
                    column=offset + 8,
                    value=requisicao["guias__insucessos__criado_em"].strftime(
                        "%H:%M:%S"
                    ),
                )
            ws.cell(
                row=ind,
                column=offset + 9,
                value=requisicao["guias__escola__codigo_eol"],
            )
            ws.cell(
                row=ind, column=offset + 10, value=requisicao["guias__codigo_unidade"]
            )
            ws.cell(
                row=ind, column=offset + 11, value=requisicao["guias__nome_unidade"]
            )
            ws.cell(
                row=ind, column=offset + 12, value=requisicao["guias__endereco_unidade"]
            )
            ws.cell(
                row=ind, column=offset + 13, value=requisicao["guias__numero_unidade"]
            )
            ws.cell(
                row=ind, column=offset + 14, value=requisicao["guias__bairro_unidade"]
            )
            ws.cell(row=ind, column=offset + 15, value=requisicao["guias__cep_unidade"])
            ws.cell(
                row=ind, column=offset + 16, value=requisicao["guias__cidade_unidade"]
            )
            ws.cell(
                row=ind, column=offset + 17, value=requisicao["guias__estado_unidade"]
            )
            ws.cell(
                row=ind, column=offset + 18, value=requisicao["guias__contato_unidade"]
            )
            ws.cell(
                row=ind, column=offset + 19, value=requisicao["guias__telefone_unidade"]
            )
            ws.cell(
                row=ind,
                column=offset + 20,
                value=requisicao["guias__alimentos__nome_alimento"],
            )
            ws.cell(
                row=ind,
                column=offset + 21,
                value=requisicao["guias__alimentos__codigo_suprimento"],
            )
            ws.cell(
                row=ind,
                column=offset + 22,
                value=requisicao["guias__alimentos__codigo_papa"],
            )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(
                    row=ind,
                    column=offset + 23,
                    value=requisicao[
                        "guias__alimentos__embalagens__descricao_embalagem"
                    ],
                )
                ws.cell(
                    row=ind,
                    column=offset + 24,
                    value=requisicao[
                        "guias__alimentos__embalagens__capacidade_embalagem"
                    ],
                )
                ws.cell(
                    row=ind,
                    column=offset + 25,
                    value=requisicao["guias__alimentos__embalagens__unidade_medida"],
                )
                ws.cell(
                    row=ind,
                    column=offset + 26,
                    value=requisicao["guias__alimentos__embalagens__qtd_volume"],
                )
            else:
                ws.cell(
                    row=ind,
                    column=offset + 27,
                    value=requisicao[
                        "guias__alimentos__embalagens__descricao_embalagem"
                    ],
                )
                ws.cell(
                    row=ind,
                    column=offset + 28,
                    value=requisicao[
                        "guias__alimentos__embalagens__capacidade_embalagem"
                    ],
                )
                ws.cell(
                    row=ind,
                    column=offset + 29,
                    value=requisicao["guias__alimentos__embalagens__unidade_medida"],
                )
                ws.cell(
                    row=ind,
                    column=offset + 30,
                    value=requisicao["guias__alimentos__embalagens__qtd_volume"],
                )
            if requisicao["guias__insucessos__criado_em"] is not None:
                ws.cell(
                    row=ind,
                    column=offset + 31,
                    value=requisicao["guias__insucessos__hora_tentativa"].strftime(
                        "%H:%M:%S"
                    ),
                )
                ws.cell(
                    row=ind,
                    column=offset + 32,
                    value=retorna_motivo_insucesso(
                        requisicao["guias__insucessos__motivo"]
                    ),
                )
                ws.cell(
                    row=ind,
                    column=offset + 33,
                    value=requisicao["guias__insucessos__justificativa"],
                )
                ws.cell(
                    row=ind,
                    column=offset + 34,
                    value=retorna_status_guia_remessa(requisicao["guias__status"]),
                )
                ws.cell(
                    row=ind,
                    column=offset + 35,
                    value=requisicao["guias__insucessos__criado_por__nome"],
                )
                ws.cell(
                    row=ind,
                    column=offset + 36,
                    value=requisicao["guias__insucessos__criado_por__cpf"],
                )

        cls.aplicar_estilo_padrao(ws, count_data, count_fields)

    @classmethod
    def cria_aba_conferencia_dilog(cls, ws, requisicoes):  # noqa C901
        cabecalho = [
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_DE_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_ALIMENTO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_DISTRIBUIDOR.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_DA_GUIA_DE_REMESSA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DA_GUIA_DE_REMESSA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.TELEFONE_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.ENDERECO_DA_UE_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.BAIRRO_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CONTATO_DA_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_EOL.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_PAPA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_CODAE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_E_HORA_DO_RECEBIMENTO_1A_CONFERENCIA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_E_HORA_DO_REGISTRO_1A_CONFERENCIA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_PREVISTA_EMBALAGEM_FECHADA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_PREVISTA_EMBALAGEM_FRACIONADA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_RECEBIDA_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_RECEBIDA_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_COMPLETO_DO_CONFERENTE_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_MOTORISTA_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.PLACA_DO_VEICULO_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DE_RECEBIMENTO_DO_ALIMENTO_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OCORRENCIAS_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OBSERVACOES_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_A_RECEBER_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_A_RECEBER_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_E_HORA_DO_RECEBIMENTO_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_E_HORA_DE_REGISTRO_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DE_RECEBIMENTO_DO_ALIMENTO_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_REPOSTA_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_REPOSTA_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_COMPLETO_DO_CONFERENTE_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_MOTORISTA_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.PLACA_DO_VEICULO_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OCORRENCIAS_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OBSERVACOES_REPOSICAO.value,
        ]

        count_fields = len(cabecalho)
        count_data = requisicoes.count()

        ws.title = "Relatório de Conferência"
        for ind, title in enumerate(cabecalho, 1):
            celula = ws.cell(row=1, column=ind)
            celula.value = title
            celula.font = Font(size="13", bold=True, color="00FFFFFF")

        for ind, requisicao in enumerate(requisicoes, 2):
            qtd_recebida = 0

            if (
                "conferencia_alimento" in requisicao
                and requisicao["conferencia_alimento"] is not None
            ):
                qtd_recebida = requisicao["conferencia_alimento"].qtd_recebido
            elif "primeira_conferencia" in requisicao:
                qtd_recebida = requisicao["guias__alimentos__embalagens__qtd_volume"]
            else:
                qtd_recebida = ""

            ws.cell(row=ind, column=1, value=requisicao["guias__data_entrega"])
            ws.cell(
                row=ind, column=2, value=requisicao["guias__alimentos__nome_alimento"]
            )
            ws.cell(row=ind, column=3, value=requisicao["distribuidor__nome_fantasia"])
            ws.cell(row=ind, column=4, value=requisicao["guias__numero_guia"])
            ws.cell(
                row=ind,
                column=5,
                value=retorna_status_guia_remessa(requisicao["guias__status"]),
            )
            ws.cell(row=ind, column=6, value=requisicao["guias__nome_unidade"])
            ws.cell(row=ind, column=7, value=requisicao["guias__telefone_unidade"])
            ws.cell(
                row=ind,
                column=8,
                value=f'{requisicao["guias__endereco_unidade"]} '
                f'{requisicao["guias__numero_unidade"]}',
            )
            ws.cell(row=ind, column=9, value=requisicao["guias__bairro_unidade"])
            ws.cell(row=ind, column=10, value=requisicao["guias__contato_unidade"])
            ws.cell(row=ind, column=11, value=requisicao["guias__escola__codigo_eol"])
            ws.cell(
                row=ind, column=12, value=requisicao["guias__alimentos__codigo_papa"]
            )
            ws.cell(row=ind, column=13, value=requisicao["guias__codigo_unidade"])
            if "primeira_conferencia" in requisicao:
                ws.cell(
                    row=ind,
                    column=14,
                    value=f'{requisicao["primeira_conferencia"].data_recebimento} '
                    f'{requisicao["primeira_conferencia"].hora_recebimento}',
                )
                ws.cell(
                    row=ind,
                    column=15,
                    value=requisicao["primeira_conferencia"].criado_em.strftime(
                        "%d/%m/%Y %H:%M:%S"
                    ),
                )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(
                    row=ind,
                    column=16,
                    value=f'{requisicao["guias__alimentos__embalagens__qtd_volume"]} '
                    f'{requisicao["guias__alimentos__embalagens__descricao_embalagem"]}',
                )
            if (
                requisicao["guias__alimentos__embalagens__tipo_embalagem"]
                == "FRACIONADA"
            ):
                ws.cell(
                    row=ind,
                    column=17,
                    value=f'{requisicao["guias__alimentos__embalagens__qtd_volume"]} '
                    f'{requisicao["guias__alimentos__embalagens__descricao_embalagem"]}',
                )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(row=ind, column=18, value=qtd_recebida)
            if (
                requisicao["guias__alimentos__embalagens__tipo_embalagem"]
                == "FRACIONADA"
            ):
                ws.cell(row=ind, column=19, value=qtd_recebida)
            if "primeira_conferencia" in requisicao:
                ws.cell(
                    row=ind,
                    column=20,
                    value=requisicao["primeira_conferencia"].criado_por.nome,
                )
                ws.cell(
                    row=ind,
                    column=21,
                    value=requisicao["primeira_conferencia"].nome_motorista,
                )
                ws.cell(
                    row=ind,
                    column=22,
                    value=requisicao["primeira_conferencia"].placa_veiculo,
                )
                if "conferencia_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=23,
                        value=retorna_status_alimento(
                            requisicao["conferencia_alimento"].status_alimento
                        ),
                    )
                    ws.cell(
                        row=ind,
                        column=24,
                        value=retorna_ocorrencias_alimento(
                            requisicao["conferencia_alimento"].ocorrencia
                        ),
                    )
                    ws.cell(
                        row=ind,
                        column=25,
                        value=requisicao["conferencia_alimento"].observacao,
                    )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(
                    row=ind,
                    column=26,
                    value=requisicao["guias__alimentos__embalagens__qtd_a_receber"],
                )
            if (
                requisicao["guias__alimentos__embalagens__tipo_embalagem"]
                == "FRACIONADA"
            ):
                ws.cell(
                    row=ind,
                    column=27,
                    value=requisicao["guias__alimentos__embalagens__qtd_a_receber"],
                )
            if "primeira_reposicao" in requisicao:
                ws.cell(
                    row=ind,
                    column=28,
                    value=f'{requisicao["primeira_reposicao"].data_recebimento} '
                    f'{requisicao["primeira_reposicao"].hora_recebimento}',
                )
                ws.cell(
                    row=ind,
                    column=29,
                    value=requisicao["primeira_reposicao"].criado_em.strftime(
                        "%d/%m/%Y %H:%M:%S"
                    ),
                )
            if "primeira_reposicao" in requisicao:
                if "reposicao_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=30,
                        value=retorna_status_alimento(
                            requisicao["reposicao_alimento"].status_alimento
                        ),
                    )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                if "reposicao_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=31,
                        value=requisicao["reposicao_alimento"].qtd_recebido,
                    )
            if (
                requisicao["guias__alimentos__embalagens__tipo_embalagem"]
                == "FRACIONADA"
            ):
                if "reposicao_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=32,
                        value=requisicao["reposicao_alimento"].qtd_recebido,
                    )
            if "primeira_reposicao" in requisicao:
                ws.cell(
                    row=ind,
                    column=33,
                    value=requisicao["primeira_reposicao"].criado_por.nome,
                )
                ws.cell(
                    row=ind,
                    column=34,
                    value=requisicao["primeira_reposicao"].nome_motorista,
                )
                ws.cell(
                    row=ind,
                    column=35,
                    value=requisicao["primeira_reposicao"].placa_veiculo,
                )
                if "reposicao_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=36,
                        value=retorna_ocorrencias_alimento(
                            requisicao["reposicao_alimento"].ocorrencia
                        ),
                    )
                    ws.cell(
                        row=ind,
                        column=37,
                        value=requisicao["reposicao_alimento"].observacao,
                    )

        cls.aplicar_estilo_padrao(ws, count_data, count_fields)

    @classmethod
    def cria_aba_conferencia_dre(cls, ws, requisicoes):  # noqa C901
        offset = 0
        cabecalho = [
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_EOL.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_CODAE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.TELEFONE_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_DE_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_DA_GUIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DA_GUIA_DE_REMESSA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_E_HORA_DO_RECEBIMENTO_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_E_HORA_DO_REGISTRO_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_E_HORA_DO_RECEBIMENTO_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_E_HORA_DO_REGISTRO_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_ALIMENTO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_PREVISTA_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_PREVISTA_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_A_REPOR_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_A_REPOR_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DE_RECEBIMENTO_DO_ALIMENTO_1A_CONFERENCIA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OCORRENCIAS_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OBSERVACOES_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_COMPLETO_DO_CONFERENTE_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_REPOSTA_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_REPOSTA_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DE_RECEBIMENTO_DO_ALIMENTO_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OCORRENCIAS_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OBSERVACOES_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_COMPLETO_DO_CONFERENTE_REPOSICAO.value,
        ]

        offset = 1

        count_fields = len(cabecalho)
        count_data = requisicoes.count()

        ws.title = "Relatório de Conferência"
        for ind, title in enumerate(cabecalho, 1):
            celula = ws.cell(row=1, column=ind)
            celula.value = title
            celula.font = Font(size="13", bold=True, color="00FFFFFF")

        for ind, requisicao in enumerate(requisicoes, 2):
            qtd_recebida = 0

            if (
                "conferencia_alimento" in requisicao
                and requisicao["conferencia_alimento"] is not None
            ):
                qtd_recebida = requisicao["conferencia_alimento"].qtd_recebido
            elif "primeira_conferencia" in requisicao:
                qtd_recebida = requisicao["guias__alimentos__embalagens__qtd_volume"]
            else:
                qtd_recebida = ""

            ws.cell(
                row=ind, column=offset, value=requisicao["guias__escola__codigo_eol"]
            )
            ws.cell(
                row=ind, column=offset + 1, value=requisicao["guias__codigo_unidade"]
            )
            ws.cell(row=ind, column=offset + 2, value=requisicao["guias__nome_unidade"])
            ws.cell(
                row=ind, column=offset + 3, value=requisicao["guias__telefone_unidade"]
            )
            ws.cell(row=ind, column=offset + 4, value=requisicao["guias__data_entrega"])
            ws.cell(row=ind, column=offset + 5, value=requisicao["guias__numero_guia"])
            ws.cell(
                row=ind,
                column=offset + 6,
                value=retorna_status_guia_remessa(requisicao["guias__status"]),
            )
            if "primeira_conferencia" in requisicao:
                ws.cell(
                    row=ind,
                    column=offset + 7,
                    value=f'{requisicao["primeira_conferencia"].data_recebimento} '
                    f'{requisicao["primeira_conferencia"].hora_recebimento}',
                )
                ws.cell(
                    row=ind,
                    column=offset + 8,
                    value=f'{requisicao["primeira_conferencia"].criado_em.strftime(FORMATO_DATA_BRASILEIRO)} '
                    f'{requisicao["primeira_conferencia"].criado_em.strftime("%H:%M:%S")}',
                )
            if "primeira_reposicao" in requisicao:
                ws.cell(
                    row=ind,
                    column=offset + 9,
                    value=f'{requisicao["primeira_reposicao"].data_recebimento} '
                    f'{requisicao["primeira_reposicao"].hora_recebimento}',
                )
                ws.cell(
                    row=ind,
                    column=offset + 10,
                    value=f'{requisicao["primeira_reposicao"].criado_em.strftime(FORMATO_DATA_BRASILEIRO)} '
                    f'{requisicao["primeira_reposicao"].criado_em.strftime("%H:%M:%S")}',
                )
            ws.cell(
                row=ind,
                column=offset + 11,
                value=requisicao["guias__alimentos__nome_alimento"],
            )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(
                    row=ind,
                    column=offset + 12,
                    value=f'{requisicao["guias__alimentos__embalagens__qtd_volume"]} '
                    f'{requisicao["guias__alimentos__embalagens__descricao_embalagem"]}',
                )
            else:
                ws.cell(
                    row=ind,
                    column=offset + 13,
                    value=f'{requisicao["guias__alimentos__embalagens__qtd_volume"]} '
                    f'{requisicao["guias__alimentos__embalagens__descricao_embalagem"]}',
                )

            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(
                    row=ind,
                    column=offset + 14,
                    value=requisicao["guias__alimentos__embalagens__qtd_a_receber"],
                )
            else:
                ws.cell(
                    row=ind,
                    column=offset + 15,
                    value=requisicao["guias__alimentos__embalagens__qtd_a_receber"],
                )
            if "primeira_conferencia" in requisicao:
                if "conferencia_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=offset + 16,
                        value=retorna_status_alimento(
                            requisicao["conferencia_alimento"].status_alimento
                        ),
                    )
                    ws.cell(
                        row=ind,
                        column=offset + 17,
                        value=retorna_ocorrencias_alimento(
                            requisicao["conferencia_alimento"].ocorrencia
                        ),
                    )
                    ws.cell(
                        row=ind,
                        column=offset + 18,
                        value=requisicao["conferencia_alimento"].observacao,
                    )
                    ws.cell(
                        row=ind,
                        column=offset + 19,
                        value=requisicao["primeira_conferencia"].criado_por.nome,
                    )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(row=ind, column=offset + 20, value=qtd_recebida)
            else:
                ws.cell(row=ind, column=offset + 21, value=qtd_recebida)
            if "primeira_reposicao" in requisicao:
                if "reposicao_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=offset + 22,
                        value=retorna_status_alimento(
                            requisicao["reposicao_alimento"].status_alimento
                        ),
                    )
                    ws.cell(
                        row=ind,
                        column=offset + 23,
                        value=retorna_ocorrencias_alimento(
                            requisicao["reposicao_alimento"].ocorrencia
                        ),
                    )
                    ws.cell(
                        row=ind,
                        column=offset + 24,
                        value=requisicao["reposicao_alimento"].observacao,
                    )
                    ws.cell(
                        row=ind,
                        column=offset + 25,
                        value=requisicao["primeira_reposicao"].criado_por.nome,
                    )
        cls.aplicar_estilo_padrao(ws, count_data, count_fields)

    @classmethod
    def cria_aba_conferencia_distribuidor(cls, ws, requisicoes):  # noqa C901
        offset = 0
        cabecalho = [
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_DA_REQUISICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_DE_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_ALIMENTO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_SUPRI.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NUMERO_DA_GUIA_DE_REMESSA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DA_GUIA_DE_REMESSA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CODIGO_CODAE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.AGRUP.value,
            StringsCabecalhoXLSXGuiaDaRemessa.TELEFONE_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.ENDERECO_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.BAIRRO_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CEP_DA_UE.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CONTATO_DA_ENTREGA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_PREVISTA_EMBALAGEM_FECHADA_3.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CAPACIDADE_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_PREVISTA_EMBALAGEM_FRACIONADA_3.value,
            StringsCabecalhoXLSXGuiaDaRemessa.CAPACIDADE_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_E_HORA_DO_RECEBIMENTO_1A_CONFERENCIA_3.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_COMPLETO_DO_CONFERENTE_1A_CONFERENCIA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_RECEBIDA_EMBALAGEM_FECHADA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_A_REPOR_REFERENTE_A_QUANTIDADE_A_RECEBER_EMBALAGEM_FECHADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_RECEBIDA_EMBALAGEM_FRACIONADA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_A_REPOR_REFERENTE_A_QUANTIDADE_A_RECEBER_EMBALAGEM_FRACIONADA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_MOTORISTA_1A_CONFERENCIA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.PLACA_DO_VEICULO_1A_CONFERENCIA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OCORRENCIAS_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OBSERVACAO_1A_CONFERENCIA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.REPOSICAO_DATA_E_HORA_DO_RECEBIMENTO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.DATA_DE_REGISTRO_DA_REPOSICAO_COM_HORA.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_COMPLETO_DO_CONFERENTE_REPOSICAO_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_REPOSTA_EMBALAGEM_FECHADA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.QUANTIDADE_REPOSTA_EMBALAGEM_FRACIONADA_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.NOME_DO_MOTORISTA_REPOSICAO_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.PLACA_DO_VEICULO_REPOSICAO_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.STATUS_DE_RECEBIMENTO_DO_ALIMENTO_REPOSICAO_2.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OCORRENCIAS_REPOSICAO.value,
            StringsCabecalhoXLSXGuiaDaRemessa.OBSERVACOES_REPOSICAO.value,
        ]

        count_fields = len(cabecalho)
        count_data = requisicoes.count()

        ws.title = "Relatório de Conferência"
        for ind, title in enumerate(cabecalho, 1):
            celula = ws.cell(row=1, column=ind)
            celula.value = title
            celula.font = Font(size="13", bold=True, color="00FFFFFF")

        for ind, requisicao in enumerate(requisicoes, 2):
            if (
                "conferencia_alimento" in requisicao
                and requisicao["conferencia_alimento"] is not None
            ):
                qtd_recebida = requisicao["conferencia_alimento"].qtd_recebido
            elif "primeira_conferencia" in requisicao:
                qtd_recebida = requisicao["guias__alimentos__embalagens__qtd_volume"]
            else:
                qtd_recebida = ""

            ws.cell(row=ind, column=offset + 1, value=requisicao["numero_solicitacao"])
            ws.cell(row=ind, column=offset + 2, value=requisicao["guias__data_entrega"])
            ws.cell(
                row=ind,
                column=offset + 3,
                value=requisicao["guias__alimentos__nome_alimento"],
            )
            ws.cell(
                row=ind,
                column=offset + 4,
                value=requisicao["guias__alimentos__codigo_suprimento"],
            )
            ws.cell(row=ind, column=offset + 5, value=requisicao["guias__numero_guia"])
            ws.cell(
                row=ind,
                column=offset + 6,
                value=retorna_status_guia_remessa(requisicao["guias__status"]),
            )
            ws.cell(row=ind, column=offset + 7, value=requisicao["guias__nome_unidade"])
            ws.cell(
                row=ind, column=offset + 8, value=requisicao["guias__codigo_unidade"]
            )
            ws.cell(
                row=ind,
                column=offset + 9,
                value=requisicao["guias__escola__subprefeitura__agrupamento"],
            )
            ws.cell(
                row=ind, column=offset + 10, value=requisicao["guias__telefone_unidade"]
            )
            ws.cell(
                row=ind,
                column=offset + 11,
                value=f'{requisicao["guias__endereco_unidade"]} '
                f'{requisicao["guias__numero_unidade"]}',
            )
            ws.cell(
                row=ind, column=offset + 12, value=requisicao["guias__bairro_unidade"]
            )
            ws.cell(row=ind, column=offset + 13, value=requisicao["guias__cep_unidade"])
            ws.cell(
                row=ind, column=offset + 14, value=requisicao["guias__contato_unidade"]
            )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(
                    row=ind,
                    column=offset + 15,
                    value=f'{requisicao["guias__alimentos__embalagens__qtd_volume"]} '
                    f'{requisicao["guias__alimentos__embalagens__descricao_embalagem"]}',
                )
                ws.cell(
                    row=ind,
                    column=offset + 16,
                    value=f'{requisicao["guias__alimentos__embalagens__capacidade_embalagem"]} '
                    f'{requisicao["guias__alimentos__embalagens__unidade_medida"]}',
                )
            if (
                requisicao["guias__alimentos__embalagens__tipo_embalagem"]
                == "FRACIONADA"
            ):
                ws.cell(
                    row=ind,
                    column=offset + 17,
                    value=f'{requisicao["guias__alimentos__embalagens__qtd_volume"]} '
                    f'{requisicao["guias__alimentos__embalagens__descricao_embalagem"]}',
                )
                ws.cell(
                    row=ind,
                    column=offset + 18,
                    value=f'{requisicao["guias__alimentos__embalagens__capacidade_embalagem"]} '
                    f'{requisicao["guias__alimentos__embalagens__unidade_medida"]}',
                )
            if "primeira_conferencia" in requisicao:
                ws.cell(
                    row=ind,
                    column=offset + 19,
                    value=f'{requisicao["primeira_conferencia"].data_recebimento} '
                    f'{requisicao["primeira_conferencia"].hora_recebimento}',
                )
                ws.cell(
                    row=ind,
                    column=offset + 20,
                    value=requisicao["primeira_conferencia"].criado_por.nome,
                )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                ws.cell(row=ind, column=offset + 21, value=qtd_recebida)
                ws.cell(
                    row=ind,
                    column=offset + 22,
                    value=requisicao["guias__alimentos__embalagens__qtd_a_receber"],
                )
            if (
                requisicao["guias__alimentos__embalagens__tipo_embalagem"]
                == "FRACIONADA"
            ):
                ws.cell(row=ind, column=offset + 23, value=qtd_recebida)
                ws.cell(
                    row=ind,
                    column=offset + 24,
                    value=requisicao["guias__alimentos__embalagens__qtd_a_receber"],
                )
            if "primeira_conferencia" in requisicao:
                ws.cell(
                    row=ind,
                    column=offset + 25,
                    value=requisicao["primeira_conferencia"].nome_motorista,
                )
                ws.cell(
                    row=ind,
                    column=offset + 26,
                    value=requisicao["primeira_conferencia"].placa_veiculo,
                )
                if "conferencia_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=offset + 27,
                        value=retorna_ocorrencias_alimento(
                            requisicao["conferencia_alimento"].ocorrencia
                        ),
                    )
                    ws.cell(
                        row=ind,
                        column=offset + 28,
                        value=requisicao["conferencia_alimento"].observacao,
                    )
            if "primeira_reposicao" in requisicao:
                ws.cell(
                    row=ind,
                    column=offset + 29,
                    value=f'{requisicao["primeira_reposicao"].data_recebimento} '
                    f'{requisicao["primeira_reposicao"].hora_recebimento}',
                )
                ws.cell(
                    row=ind,
                    column=offset + 30,
                    value=f'{requisicao["primeira_reposicao"].criado_em.strftime(FORMATO_DATA_BRASILEIRO)} '
                    f'{requisicao["primeira_reposicao"].criado_em.strftime("%H:%M:%S")}',
                )
            if "primeira_reposicao" in requisicao:
                if "reposicao_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=offset + 31,
                        value=requisicao["primeira_reposicao"].criado_por.nome,
                    )
            if requisicao["guias__alimentos__embalagens__tipo_embalagem"] == "FECHADA":
                if "reposicao_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=offset + 32,
                        value=requisicao["reposicao_alimento"].qtd_recebido,
                    )
            if (
                requisicao["guias__alimentos__embalagens__tipo_embalagem"]
                == "FRACIONADA"
            ):
                if "reposicao_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=offset + 33,
                        value=requisicao["reposicao_alimento"].qtd_recebido,
                    )
            if "primeira_reposicao" in requisicao:
                ws.cell(
                    row=ind,
                    column=offset + 34,
                    value=requisicao["primeira_reposicao"].nome_motorista,
                )
                ws.cell(
                    row=ind,
                    column=offset + 35,
                    value=requisicao["primeira_reposicao"].placa_veiculo,
                )
                if "reposicao_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=offset + 36,
                        value=retorna_status_alimento(
                            requisicao["reposicao_alimento"].status_alimento
                        ),
                    )
                if "reposicao_alimento" in requisicao:
                    ws.cell(
                        row=ind,
                        column=offset + 37,
                        value=retorna_ocorrencias_alimento(
                            requisicao["reposicao_alimento"].ocorrencia
                        ),
                    )
                    ws.cell(
                        row=ind,
                        column=offset + 38,
                        value=requisicao["reposicao_alimento"].observacao,
                    )

        cls.aplicar_estilo_padrao(ws, count_data, count_fields)

    @classmethod  # noqa C901
    def exportar_entregas(
        cls,
        requisicoes,
        requisicoes_insucesso,
        perfil,
        tem_conferencia,
        tem_insucesso,
        is_async=False,
    ):
        wb = Workbook()
        ws_conferencia = wb.active
        if perfil == "DISTRIBUIDOR":
            cls.cria_aba_conferencia_distribuidor(ws_conferencia, requisicoes)
        elif perfil == "DRE":
            cls.cria_aba_conferencia_dre(ws_conferencia, requisicoes)
        else:
            cls.cria_aba_conferencia_dilog(ws_conferencia, requisicoes)

        if tem_conferencia and not tem_insucesso:
            pass
        elif tem_insucesso and not tem_conferencia:
            ws_insucesso = wb.active
            cls.cria_aba_insucesso(ws_insucesso, requisicoes_insucesso, perfil)
        else:
            ws_insucesso = wb.create_sheet("Relatório de Insucesso")
            cls.cria_aba_insucesso(ws_insucesso, requisicoes_insucesso, perfil)

        arquivo = cls.gera_arquivo(wb)

        return arquivo if is_async else {"arquivo": arquivo}
