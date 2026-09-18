from io import BytesIO

import openpyxl
import pytest
from rest_framework import status

from src.dados_comuns.models import CentralDeDownload
from src.pre_recebimento.cronograma_semanal.api.relatorio_cronograma_semanal_excel import (
    COLUNAS,
    MENSAGEM_SEM_REGISTROS,
    _sanitiza_texto,
    gera_relatorio_cronogramas_semanais_xlsx,
)
from src.pre_recebimento.tasks import (
    gerar_relatorio_cronogramas_semanais_xlsx_async,
)

pytestmark = pytest.mark.django_db

NOME_ABA = "Cronogramas Semanais"
URL_EXPORTAR = "/cronogramas-semanais/gerar-relatorio-xlsx-async/"
LINHA_CABECALHO = 1
PRIMEIRA_LINHA_DE_DADOS = 2


def _abre_planilha(arquivo):
    assert isinstance(arquivo, bytes)
    return openpyxl.load_workbook(filename=BytesIO(arquivo))[NOME_ABA]


def _linhas_de_dados(sheet):
    """Linhas preenchidas abaixo do cabeçalho."""
    return [
        linha
        for linha in sheet.iter_rows(min_row=PRIMEIRA_LINHA_DE_DADOS, values_only=True)
        if any(valor is not None for valor in linha)
    ]


def test_planilha_comeca_pelo_cabecalho_das_colunas(
    cronograma_semanal_completo_para_excel,
):
    """Sem título nem subtítulo: a primeira linha já é o cabeçalho."""
    arquivo = gera_relatorio_cronogramas_semanais_xlsx(
        [cronograma_semanal_completo_para_excel.id]
    )
    sheet = _abre_planilha(arquivo)

    cabecalho = tuple(
        sheet.iter_rows(
            min_row=LINHA_CABECALHO, max_row=LINHA_CABECALHO, values_only=True
        )
    )[0]
    assert cabecalho == tuple(COLUNAS)


def test_conteudo_da_linha(cronograma_semanal_completo_para_excel):
    """Cada coluna traz o dado da sua origem: semanal, mensal ou programação."""
    arquivo = gera_relatorio_cronogramas_semanais_xlsx(
        [cronograma_semanal_completo_para_excel.id]
    )
    sheet = _abre_planilha(arquivo)

    linha = [
        sheet.cell(row=PRIMEIRA_LINHA_DE_DADOS, column=coluna).value
        for coluna in range(1, len(COLUNAS) + 1)
    ]

    assert linha[0] == "012/2026A"  # Nº do Cronograma Mensal
    assert linha[1] == "012/2026P"  # Nº do Cronograma Semanal
    assert linha[2] == "BELA VISTA"
    assert linha[3] == "CAQUI"
    assert linha[4] == pytest.approx(25000.0)
    assert linha[5] == "un"
    assert linha[6] == pytest.approx(1.35)
    assert linha[7] == "20/06/2026"  # Período Programado Inicial
    assert linha[8] == "23/06/2026"  # Período Programado Final
    assert linha[9] == pytest.approx(8250.0)
    assert linha[10] == "un"
    assert linha[11] == "444.2026/004444-4"
    assert linha[12] == "44.444/2026"
    assert linha[13] == "44/SME/CODAE/2026"
    assert linha[14] == "Pregão Eletronico"
    assert linha[15] == "Enviado ao Fornecedor"


def test_cada_programacao_vira_uma_linha(cronograma_semanal_completo_para_excel):
    """O cronograma tem duas programações, então ocupa duas linhas, com os
    dados do cronograma repetidos."""
    arquivo = gera_relatorio_cronogramas_semanais_xlsx(
        [cronograma_semanal_completo_para_excel.id]
    )
    sheet = _abre_planilha(arquivo)

    linhas = _linhas_de_dados(sheet)

    assert len(linhas) == 2
    assert [linha[7] for linha in linhas] == ["20/06/2026", "25/06/2026"]
    assert {linha[1] for linha in linhas} == {"012/2026P"}


def test_filtro_de_mes_recorta_as_linhas(
    cronograma_semanal_completo_para_excel, programacao_de_julho
):
    """Com três programações (duas em junho, uma em julho), filtrar julho
    deixa só a linha de julho."""
    arquivo = gera_relatorio_cronogramas_semanais_xlsx(
        [cronograma_semanal_completo_para_excel.id],
        {"mes_inicial": "07/2026", "mes_final": "07/2026"},
    )
    sheet = _abre_planilha(arquivo)

    linhas = _linhas_de_dados(sheet)

    assert len(linhas) == 1
    assert linhas[0][7] == "01/07/2026"
    assert linhas[0][8] == "10/07/2026"


def test_cronograma_sem_programacao_nao_gera_linha(
    cronograma_semanal_sem_programacoes,
):

    arquivo = gera_relatorio_cronogramas_semanais_xlsx(
        [cronograma_semanal_sem_programacoes.id]
    )
    sheet = _abre_planilha(arquivo)

    assert sheet.cell(row=PRIMEIRA_LINHA_DE_DADOS, column=1).value == (
        MENSAGEM_SEM_REGISTROS
    )


def test_colunas_numericas_sao_numeros_formatados(
    cronograma_semanal_completo_para_excel,
):
    """Quantidades e custo precisam ser número no Excel, não texto."""
    arquivo = gera_relatorio_cronogramas_semanais_xlsx(
        [cronograma_semanal_completo_para_excel.id]
    )
    sheet = _abre_planilha(arquivo)

    assert sheet.cell(row=PRIMEIRA_LINHA_DE_DADOS, column=5).number_format == (
        "#,##0.00"
    )
    assert sheet.cell(row=PRIMEIRA_LINHA_DE_DADOS, column=10).number_format == (
        "#,##0.00"
    )
    assert sheet.cell(row=PRIMEIRA_LINHA_DE_DADOS, column=7).number_format == (
        "R$ #,##0.00"
    )


def test_planilha_sem_registros():
    arquivo = gera_relatorio_cronogramas_semanais_xlsx([])
    sheet = _abre_planilha(arquivo)

    cabecalho = tuple(
        sheet.iter_rows(
            min_row=LINHA_CABECALHO, max_row=LINHA_CABECALHO, values_only=True
        )
    )[0]
    assert cabecalho == tuple(COLUNAS)
    assert sheet.cell(row=PRIMEIRA_LINHA_DE_DADOS, column=1).value == (
        MENSAGEM_SEM_REGISTROS
    )


@pytest.mark.parametrize(
    "valor,esperado",
    [
        ("=1+1", "'=1+1"),
        ("@cmd", "'@cmd"),
        ("+1", "'+1"),
        ("-1", "'-1"),
        ("CAQUI", "CAQUI"),
        (None, ""),
    ],
)
def test_sanitiza_injecao_de_formula(valor, esperado):
    assert _sanitiza_texto(valor) == esperado


def test_task_grava_arquivo_na_central_de_downloads(
    client_autenticado_vinculo_dilog_cronograma,
    cronograma_semanal_completo_para_excel,
):
    _, usuario = client_autenticado_vinculo_dilog_cronograma

    gerar_relatorio_cronogramas_semanais_xlsx_async(
        usuario.username,
        [cronograma_semanal_completo_para_excel.id],
        None,
    )

    assert CentralDeDownload.objects.count() == 1
    arquivo = CentralDeDownload.objects.first()
    assert arquivo.identificador == "relatorio_cronogramas_semanais.xlsx"
    assert arquivo.status == CentralDeDownload.STATUS_CONCLUIDO
    assert arquivo.arquivo.size > 0


def test_endpoint_enfileira_com_ids_filtrados_e_mes(
    client_autenticado_vinculo_dilog_cronograma,
    cronograma_semanal_completo_para_excel,
    cronograma_semanal_sem_programacoes,
    monkeypatch,
):
    """O endpoint aplica os filtros da tela e repassa só os ids e o mês."""
    client, usuario = client_autenticado_vinculo_dilog_cronograma
    chamadas = []

    monkeypatch.setattr(
        "src.pre_recebimento.cronograma_semanal.api.viewsets"
        ".gerar_relatorio_cronogramas_semanais_xlsx_async.delay",
        lambda *args: chamadas.append(args),
    )

    response = client.get(
        URL_EXPORTAR,
        {"numero_cronograma_semanal": "012/2026P", "mes_inicial": "06/2026"},
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.json() == {
        "detail": "Solicitação de geração de arquivo recebida com sucesso."
    }

    username, ids_cronogramas, filtros = chamadas[0]
    assert username == usuario.username
    assert ids_cronogramas == [cronograma_semanal_completo_para_excel.id]
    assert cronograma_semanal_sem_programacoes.id not in ids_cronogramas
    assert filtros == {"mes_inicial": "06/2026", "mes_final": None}
