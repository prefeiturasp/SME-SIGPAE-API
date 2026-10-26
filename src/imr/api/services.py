from datetime import datetime
from itertools import groupby
from operator import attrgetter

from django.http import HttpResponse
from openpyxl import Workbook, styles
from openpyxl.worksheet.datavalidation import DataValidation

from src.dados_comuns.constants import FORMATO_DATA_BRASILEIRO
from src.dados_comuns.helper_planilha_modelo import (
    cria_validacao_lista_em_sheet_oculto,
)
from src.imr.models import (
    CategoriaOcorrencia,
    FormularioDiretor,
    OcorrenciaNaoSeAplica,
    TipoGravidade,
    TipoPenalidade,
    TipoPerguntaParametrizacaoOcorrencia,
)
from src.terceirizada.models import Edital


def exportar_planilha_importacao_tipos_penalidade():
    """Gera a planilha modelo de importação dos tipos de penalidade."""
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        "attachment; filename=planilha_importacao_tipos_penalidade.xlsx"
    )
    workbook: Workbook = Workbook()
    ws = workbook.active
    ws.title = "Tipos de Penalidade"
    headers = [
        "Edital",
        "Número da Cláusula/Item",
        "Gravidade",
        "Obrigações (separadas por ;)",
        "Descrição da Cláusula/Item",
        "Status",
    ]
    sheet = workbook["Tipos de Penalidade"]
    for column in "ABCDEFGF":
        sheet.column_dimensions[column].width = 30
    _font = styles.Font(name="Calibri", sz=10)
    {k: setattr(styles.DEFAULT_FONT, k, v) for k, v in _font.__dict__.items()}
    for i in range(0, len(headers)):
        cabecalho = ws.cell(row=1, column=1 + i, value=headers[i])
        cabecalho.fill = styles.PatternFill("solid", fgColor="ffff99")
        cabecalho.font = styles.Font(name="Calibri", size=10, bold=True)
        cabecalho.border = styles.Border(
            left=styles.Side(border_style="thin", color="000000"),
            right=styles.Side(border_style="thin", color="000000"),
            top=styles.Side(border_style="thin", color="000000"),
            bottom=styles.Side(border_style="thin", color="000000"),
        )

    hidden_sheet = workbook.create_sheet(title="HiddenSheet")
    hidden_sheet.sheet_state = "hidden"

    editais = [edital.numero for edital in Edital.objects.all()]
    ws = cria_validacao_lista_em_sheet_oculto(
        editais, hidden_sheet, ws, workbook, "A", 1, "Edital", "o", "A"
    )

    tipos_gravidade = ", ".join(
        [gravidade.tipo for gravidade in TipoGravidade.objects.all()]
    )
    dv2 = DataValidation(type="list", formula1=f'"{tipos_gravidade}"', allow_blank=True)
    dv2.error = "Tipo de Gravidade Inválido"
    dv2.errorTitle = "Tipo de Gravidade não permitido"
    ws.add_data_validation(dv2)
    dv2.add("C2:C1048576")

    dv3 = DataValidation(type="list", formula1='"Ativo,Inativo"', allow_blank=True)
    dv3.error = "Status Inválido"
    dv3.errorTitle = "Status não permitido"
    ws.add_data_validation(dv3)
    dv3.add("F2:F1048576")

    workbook.save(response)

    return response


def exportar_planilha_importacao_tipos_ocorrencia(request, **kwargs):
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        "attachment; filename=planilha_importacao_tipos_ocorrencia.xlsx"
    )
    workbook: Workbook = Workbook()
    ws = workbook.active
    ws.title = "Tipos de Ocorrência"
    headers = [
        "Posição",
        "Perfis",
        "Edital",
        "Categoria da Ocorrência",
        "Título",
        "Descrição",
        "Penalidade",
        "É IMR?",
        "Pontuação (IMR)",
        "Tolerância",
        "% de Desconto",
        "Status",
        "Aceita múltiplas respostas?",
    ]
    sheet = workbook["Tipos de Ocorrência"]
    for column in "ABCHIJKL":
        sheet.column_dimensions[column].width = 20
    for column in "EFM":
        sheet.column_dimensions[column].width = 35
    for column in "DG":
        sheet.column_dimensions[column].width = 50
    _font = styles.Font(name="Calibri", sz=10)
    {k: setattr(styles.DEFAULT_FONT, k, v) for k, v in _font.__dict__.items()}
    for i in range(0, len(headers)):
        cabecalho = ws.cell(row=1, column=1 + i, value=headers[i])
        cabecalho.fill = styles.PatternFill("solid", fgColor="ffff99")
        cabecalho.font = styles.Font(name="Calibri", size=10, bold=True)
        cabecalho.border = styles.Border(
            left=styles.Side(border_style="thin", color="000000"),
            right=styles.Side(border_style="thin", color="000000"),
            top=styles.Side(border_style="thin", color="000000"),
            bottom=styles.Side(border_style="thin", color="000000"),
        )

    hidden_sheet = workbook.create_sheet(title="HiddenSheet")
    hidden_sheet.sheet_state = "hidden"

    dv = DataValidation(
        type="list",
        formula1='"DIRETOR,SUPERVISAO,DIRETOR E SUPERVISAO"',
        allow_blank=True,
    )
    dv.error = "Perfil Inválido"
    dv.errorTitle = "Perfil não permitido"
    ws.add_data_validation(dv)
    dv.add("B2:B1048576")

    editais = [edital.numero for edital in Edital.objects.all()]
    ws = cria_validacao_lista_em_sheet_oculto(
        editais, hidden_sheet, ws, workbook, "A", 1, "Edital", "o", "C"
    )

    categorias = [categoria.nome for categoria in CategoriaOcorrencia.objects.all()]
    ws = cria_validacao_lista_em_sheet_oculto(
        categorias, hidden_sheet, ws, workbook, "B", 2, "Categoria", "a", "D"
    )

    penalidades = [
        f"{penalidade.edital.numero} - {penalidade.numero_clausula}"
        for penalidade in TipoPenalidade.objects.all()
    ]
    ws = cria_validacao_lista_em_sheet_oculto(
        penalidades, hidden_sheet, ws, workbook, "C", 3, "Penalidade", "a", "G"
    )

    dv4 = DataValidation(type="list", formula1='"SIM,NÃO"', allow_blank=True)
    dv4.error = "Opção Inválida"
    dv4.errorTitle = "Opção não permitida"
    ws.add_data_validation(dv4)
    dv4.add("H2:H1048576")

    dv5 = DataValidation(type="list", formula1='"Ativo,Inativo"', allow_blank=True)
    dv5.error = "Status Inválido"
    dv5.errorTitle = "Status não permitido"
    ws.add_data_validation(dv5)
    dv5.add("L2:L1048576")

    dv6 = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    dv6.error = "Valor Inválido"
    dv6.errorTitle = "Valor não permitido"
    ws.add_data_validation(dv6)
    dv6.add("M2:M1048576")

    workbook.save(response)

    return response


class RelatorioNotificacaoService:
    CATEGORIA_QUANTIDADE_QUALIDADE = (
        "QUANTIDADE/QUALIDADE DE UTENSÍLIOS/MOBILIÁRIOS/EQUIPAMENTOS"
    )

    def __init__(self, formulario_supervisao):
        categoria_reparo_e_adaptacao = "MANUTENÇÃO DE EQUIPAMENTOS/REPARO E ADAPTAÇÃO"
        self.categoria_template_map = {
            self.CATEGORIA_QUANTIDADE_QUALIDADE: (
                "IMR/relatorio_de_notificacao/quantidade_qualidade/pdf.html"
            ),
            categoria_reparo_e_adaptacao: (
                "IMR/relatorio_de_notificacao/reparo_e_adaptacao/pdf.html"
            ),
        }
        self.formulario_supervisao = formulario_supervisao

    def get_lista_respostas(self, categoria):
        respostas = self.formulario_supervisao.formulario_base.buscar_respostas(
            categoria=categoria
        )
        respostas.sort(key=attrgetter("parametrizacao.tipo_ocorrencia.uuid", "grupo"))
        grouped = groupby(
            respostas, key=attrgetter("parametrizacao.tipo_ocorrencia.uuid", "grupo")
        )
        lista_respostas = []

        for _, group in grouped:
            _respostas = []
            for resposta in group:
                _respostas.append(resposta)

            _respostas_sorted = sorted(
                _respostas, key=lambda x: int(x.parametrizacao.posicao)
            )
            lista_respostas.append(_respostas_sorted)
        return lista_respostas

    def formatar_respostas(self, categoria):
        lista_respostas = self.get_lista_respostas(categoria)

        if categoria != self.CATEGORIA_QUANTIDADE_QUALIDADE:
            return lista_respostas

        respostas_formatadas = []
        for respostas in lista_respostas:
            grupo_formatado = list(respostas)
            for indice in range(2, min(5, len(grupo_formatado))):
                grupo_formatado[indice] = int(grupo_formatado[indice].resposta)
            respostas_formatadas.append(grupo_formatado)

        return respostas_formatadas

    def retornar_dados_formatados(self, categoria):
        dados = {
            "diretoria_regional": self.formulario_supervisao.escola.diretoria_regional.nome,
            "unidade": (
                f"{self.formulario_supervisao.escola.codigo_eol} - "
                f"{self.formulario_supervisao.escola.nome}"
            ),
            "maior_frequencia_no_periodo": self.formulario_supervisao.maior_frequencia_no_periodo,
            "total_matriculados_por_data": self.formulario_supervisao.escola.quantidade_alunos_matriculados_por_data(
                self.formulario_supervisao.formulario_base.data
            ),
            "data_visita": self.formulario_supervisao.formulario_base.data.strftime(
                FORMATO_DATA_BRASILEIRO
            ),
            "usuario": self.formulario_supervisao.formulario_base.usuario,
            "lote": self.formulario_supervisao.escola.lote.nome,
            "terceirizada": self.formulario_supervisao.escola.lote.terceirizada.nome_fantasia,
            "edital": (
                f"Edital nº {self.formulario_supervisao.escola.edital.numero}"
                if self.formulario_supervisao.escola.edital
                else "-"
            ),
            "respostas": self.formatar_respostas(categoria),
            "data_geracao": datetime.now().strftime(FORMATO_DATA_BRASILEIRO),
        }
        return dados


def buscar_respostas_diretor(formularios_base_ids):
    respostas_encontradas = []
    perguntas = TipoPerguntaParametrizacaoOcorrencia.objects.select_related(
        "tipo_resposta"
    )
    for pergunta in perguntas:
        modelo_resposta = pergunta.get_model_tipo_resposta()
        respostas = modelo_resposta.objects.filter(
            formulario_base_id__in=formularios_base_ids
        ).select_related("parametrizacao__tipo_ocorrencia__categoria")
        if modelo_resposta._meta.get_field("resposta").is_relation:
            respostas = respostas.select_related("resposta")

        for resposta in respostas:
            parametrizacao = resposta.parametrizacao
            detalhe = {
                "uuid": resposta.uuid,
                "parametrizacao": parametrizacao.uuid,
                "titulo": parametrizacao.titulo,
                "posicao": parametrizacao.posicao,
                "resposta": resposta.resposta_str,
            }
            respostas_encontradas.append(
                (resposta, parametrizacao.tipo_ocorrencia, detalhe)
            )

    respostas_nao_se_aplica = OcorrenciaNaoSeAplica.objects.filter(
        formulario_base_id__in=formularios_base_ids
    ).select_related("tipo_ocorrencia__categoria")
    for resposta in respostas_nao_se_aplica:
        detalhe = {
            "uuid": resposta.uuid,
            "parametrizacao": None,
            "titulo": "Não se aplica",
            "posicao": None,
            "resposta": resposta.descricao or "Não se aplica",
        }
        respostas_encontradas.append((resposta, resposta.tipo_ocorrencia, detalhe))

    return respostas_encontradas


def _valor_resposta_para_edicao(resposta):
    if isinstance(resposta, OcorrenciaNaoSeAplica):
        return resposta.descricao

    valor = resposta.resposta

    if resposta._meta.get_field("resposta").is_relation:
        return valor.uuid

    if isinstance(valor, list):
        return [data.strftime(FORMATO_DATA_BRASILEIRO) for data in valor]

    return valor


def obter_ocorrencia_diretor(formulario, tipo_ocorrencia, grupo):
    respostas = [
        {
            **detalhe,
            "resposta": _valor_resposta_para_edicao(resposta),
        }
        for resposta, tipo, detalhe in buscar_respostas_diretor(
            [formulario.formulario_base_id]
        )
        if tipo.pk == tipo_ocorrencia.pk and resposta.grupo == grupo
    ]

    if not respostas:
        return None

    respostas.sort(key=lambda resposta: (resposta["posicao"] or 0, resposta["titulo"]))

    return {
        "uuid": formulario.uuid,
        "solicitacao_medicao_inicial": formulario.solicitacao_medicao_inicial.uuid,
        "data": formulario.formulario_base.data,
        "categoria": tipo_ocorrencia.categoria,
        "tipo_ocorrencia": tipo_ocorrencia.uuid,
        "titulo": tipo_ocorrencia.titulo,
        "descricao": tipo_ocorrencia.descricao,
        "grupo": grupo,
        "respostas": respostas,
    }


def _solicitacao_possui_ocorrencias(solicitacao_medicao_inicial):
    formularios = FormularioDiretor.objects.filter(
        solicitacao_medicao_inicial=solicitacao_medicao_inicial
    )
    return bool(listar_ocorrencias_diretor(formularios))


def excluir_ocorrencia_diretor(formulario, tipo_ocorrencia, grupo):
    form_base = formulario.formulario_base

    for resposta in form_base.buscar_respostas():
        if (
            resposta.parametrizacao.tipo_ocorrencia_id == tipo_ocorrencia.pk
            and resposta.grupo == grupo
        ):
            resposta.delete()

    OcorrenciaNaoSeAplica.objects.filter(
        formulario_base=form_base,
        tipo_ocorrencia=tipo_ocorrencia,
        grupo=grupo,
    ).delete()

    if (
        not form_base.buscar_respostas()
        and not form_base.respostas_nao_se_aplica.exists()
    ):
        form_base.delete()

    solicitacao_medicao_inicial = formulario.solicitacao_medicao_inicial

    if not _solicitacao_possui_ocorrencias(solicitacao_medicao_inicial):
        solicitacao_medicao_inicial.com_ocorrencias = False
        solicitacao_medicao_inicial.save(update_fields=["com_ocorrencias"])


def listar_ocorrencias_diretor(formularios):
    formularios_por_base = {
        formulario.formulario_base_id: formulario for formulario in formularios
    }
    if not formularios_por_base:
        return []

    ocorrencias = {}
    for resposta, tipo_ocorrencia, detalhe in buscar_respostas_diretor(
        list(formularios_por_base)
    ):
        formulario = formularios_por_base[resposta.formulario_base_id]
        chave = (formulario.pk, tipo_ocorrencia.pk, resposta.grupo)
        ocorrencia = ocorrencias.setdefault(
            chave,
            {
                "uuid": formulario.uuid,
                "data": formulario.formulario_base.data,
                "categoria": tipo_ocorrencia.categoria,
                "tipo_ocorrencia": tipo_ocorrencia.uuid,
                "titulo": tipo_ocorrencia.titulo,
                "descricao": tipo_ocorrencia.descricao,
                "grupo": resposta.grupo,
                "respostas": [],
            },
        )
        ocorrencia["respostas"].append(detalhe)

    for ocorrencia in ocorrencias.values():
        ocorrencia["respostas"].sort(
            key=lambda resposta: (resposta["posicao"] or 0, resposta["titulo"])
        )

    return sorted(
        ocorrencias.values(),
        key=lambda ocorrencia: ocorrencia["data"],
        reverse=True,
    )
